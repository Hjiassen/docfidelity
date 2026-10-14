"""独立审查缺陷的真实 DOCX 回归。"""

from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from docfidelity.contracts.errors import EngineError
from docfidelity.ooxml.doc import new_tree, parse, qn, sub, xpath
from docfidelity.ooxml.package import DocxPackage
from docfidelity.registry.candidates import inspect_template
from docfidelity.registry.freeze import freeze_template
from docfidelity.registry.selection import default_selection
from tests.fixtures.effective_template import build_effective_source

pytestmark = pytest.mark.unit


def freeze(source: Path, root: Path, *, formats: dict[str, bool] | None = None):
    selection = default_selection(inspect_template(source, import_id="review"))
    if formats is not None:
        from docfidelity.contracts.imports import FormatSelection

        selection.formats = FormatSelection(**formats)
    version = freeze_template(source, selection, root, template_id="t", version=1, name="T")
    return version, DocxPackage.open(root / "effective.docx")


@pytest.mark.parametrize("complex_field", [False, True])
def test_review_selected_toc_retains_safe_field_and_empty_formats(
    tmp_path: Path, complex_field: bool
) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source, read_only=False)
    body = xpath(package.part("word/document.xml").tree, "./w:body")[0]
    instruction = 'TOC \\o "1-3" \\h \\z \\u'
    first = new_tree("w:p")
    sub(sub(first, "w:pPr"), "w:pStyle", {"w:val": "TOC1"})
    if complex_field:
        sub(sub(first, "w:r"), "w:fldChar", {"w:fldCharType": "begin"})
        sub(sub(first, "w:r"), "w:instrText").text = instruction
        sub(sub(first, "w:r"), "w:fldChar", {"w:fldCharType": "separate"})
        sub(sub(first, "w:r"), "w:t").text = "OLD_TOC_FIRST"
        last = new_tree("w:p")
        sub(sub(last, "w:pPr"), "w:tabs")
        sub(sub(last, "w:r"), "w:t").text = "OLD_TOC_LAST"
        sub(sub(last, "w:r"), "w:fldChar", {"w:fldCharType": "end"})
        body.insert(2, last)
    else:
        field = sub(first, "w:fldSimple", {"w:instr": instruction})
        sub(sub(field, "w:r"), "w:t").text = "OLD_TOC_FIRST"
    body.insert(2, first)
    package.save(source)
    version, frozen = freeze(source, tmp_path / "bundle")
    fields = xpath(frozen.part("word/document.xml").read_tree, ".//w:fldSimple")
    assert any(f.get(qn("w:instr")) == instruction for f in fields)
    assert version.format_bindings.toc
    marker = version.format_bindings.toc[0].marker
    wrappers = xpath(
        frozen.part("word/document.xml").read_tree, f'.//w:sdt[w:sdtPr/w:tag[@w:val="{marker}"]]'
    )
    assert len(wrappers) == 1
    assert xpath(wrappers[0], ".//w:pPr/w:pStyle")[0].get(qn("w:val")) == "TOC1"
    assert b"OLD_TOC" not in frozen.part("word/document.xml").to_bytes()
    assert xpath(frozen.part("word/settings.xml").read_tree, './w:updateFields[@w:val="true"]')
    if complex_field:
        assert len(xpath(wrappers[0], ".//w:p")) == 2


def test_review_custom_heading_inherits_outline(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    parent = document.styles.add_style("ParentStyle", WD_STYLE_TYPE.PARAGRAPH)
    parent.base_style = document.styles["Heading 1"]
    child = document.styles.add_style("CustomHeading", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = parent
    document.add_paragraph("OLD_CUSTOM", child)
    document.add_paragraph("OLD_BODY")
    document.save(source)
    draft = inspect_template(source, import_id="review")
    assert next(n for n in draft.nodes if n.text_preview == "OLD_CUSTOM").heading_level == 1
    version, _ = freeze(source, tmp_path / "bundle")
    assert 1 in version.format_bindings.heading
    body = next(p for p in version.prototypes if p.prototype_id == version.format_bindings.body)
    assert not xpath(
        parse((tmp_path / "bundle" / body.path).read_bytes()),
        './w:pPr/w:pStyle[@w:val="CustomHeading"]',
    )


def test_review_numbering_scrubs_business_labels(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source, read_only=False)
    text = xpath(package.part("word/numbering.xml").tree, ".//w:lvlText")[0]
    text.set(qn("w:val"), "OLD_PRODUCT_LABEL 旧业务 (%1.%2)")
    package.save(source)
    version, frozen = freeze(source, tmp_path / "bundle")
    values = [
        e.get(qn("w:val"), "")
        for e in xpath(frozen.part("word/numbering.xml").read_tree, ".//w:lvlText")
    ]
    assert "(%1.%2)" in values
    assert not any("OLD_" in v or "旧业务" in v for v in values)
    assert version.cleanup_report.cleared_numbering_labels >= 1


def test_review_theme_dependencies_are_closed_and_cleaned(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source, read_only=False)
    theme = package.part("word/theme/theme1.xml").tree
    xpath(theme, ".//a:majorFont/a:latin")[0].set("typeface", "Review Font")
    xpath(theme, ".//a:clrScheme/a:accent1/a:srgbClr")[0].set("val", "123456")
    theme.set("name", "OLD_THEME_BUSINESS")
    extension = sub(theme, "a:extLst")
    sub(extension, "a:ext").text = "OLD_THEME_PAYLOAD"
    package.save(source)
    _, frozen = freeze(source, tmp_path / "bundle")
    resolved = frozen.resolve_single("word/document.xml", role="theme")
    assert resolved is not None and resolved.part is not None
    assert (
        xpath(resolved.part.read_tree, ".//a:majorFont/a:latin")[0].get("typeface") == "Review Font"
    )
    assert (
        xpath(resolved.part.read_tree, ".//a:clrScheme/a:accent1/a:srgbClr")[0].get("val")
        == "123456"
    )
    assert b"OLD_THEME" not in resolved.part.to_bytes()


def test_review_cover_cancel_removes_section_break(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    selection = default_selection(inspect_template(source, import_id="review"))
    selection.ranges = [r for r in selection.ranges if r.kind != "cover"]
    freeze_template(source, selection, tmp_path / "bundle", template_id="t", version=1, name="T")
    package = DocxPackage.open(tmp_path / "bundle/effective.docx")
    assert not xpath(package.part("word/document.xml").read_tree, "./w:body/w:p/w:pPr/w:sectPr")
    assert len(xpath(package.part("word/document.xml").read_tree, ".//w:sectPr")) == 1
    assert xpath(package.part("word/document.xml").read_tree, ".//w:sectPr/w:headerReference")


def test_review_selected_footnote_is_rejected_and_unselected_excluded(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source, read_only=False)
    paragraph = xpath(package.part("word/document.xml").tree, './/w:p[w:r/w:t="OLD_BODY"]')[0]
    sub(sub(paragraph, "w:r"), "w:footnoteReference", {"w:id": "1"})
    notes = new_tree("w:footnotes")
    sub(
        sub(sub(sub(notes, "w:footnote", {"w:id": "1"}), "w:p"), "w:r"), "w:t"
    ).text = "OLD_FOOTNOTE"
    package.add_xml_part(
        "word/footnotes.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        notes,
    )
    package.attach_relationship(
        "word/document.xml",
        type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
        target="footnotes.xml",
    )
    package.save(source)
    draft = inspect_template(source, import_id="review")
    assert any(not n.supported and "footnoteReference" in n.xpath for n in draft.nodes)
    selection = default_selection(draft)
    freeze_template(source, selection, tmp_path / "clean", template_id="t", version=1, name="T")
    assert not DocxPackage.open(tmp_path / "clean/effective.docx").has_part("word/footnotes.xml")
    paragraph_node = next(n for n in draft.nodes if n.text_preview == "OLD_BODY")
    from docfidelity.contracts.imports import SelectionRange

    selection.ranges.append(
        SelectionRange(
            kind="body", start_node_id=paragraph_node.node_id, end_node_id=paragraph_node.node_id
        )
    )
    with pytest.raises(EngineError, match="E_SCHEMA_INVALID"):
        freeze_template(source, selection, tmp_path / "bad", template_id="t", version=2, name="T")


@pytest.mark.parametrize("format_name", ["page_setup", "heading", "body", "numbering", "tables"])
def test_review_format_toggle_is_persisted_and_consumed(tmp_path: Path, format_name: str) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source, read_only=False)
    root = package.part("word/document.xml").tree
    xpath(root, ".//w:pgMar")[0].set(qn("w:left"), "2222")
    heading = xpath(root, './/w:p[w:r/w:t="OLD_HEADING"]')[0]
    sub(sub(heading, "w:r"), "w:rPr")
    sub(xpath(heading, "./w:r/w:rPr")[0], "w:color", {"w:val": "FF0000"})
    for table in xpath(root, ".//w:tbl"):
        sub(xpath(table, "./w:tblPr")[0], "w:tblStyle", {"w:val": "TableGrid"})
    package.save(source)
    enabled, _ = freeze(source, tmp_path / "enabled")
    disabled, frozen = freeze(source, tmp_path / "disabled", formats={format_name: False})
    assert disabled.content_hash != enabled.content_hash
    assert not getattr(disabled.selection.formats, format_name)
    if format_name == "page_setup":
        assert all(
            e.get(qn("w:left")) == "1440"
            for e in xpath(frozen.part("word/document.xml").read_tree, ".//w:pgMar")
        )
    elif format_name == "numbering":
        assert frozen.has_part("word/numbering.xml")
        assert not xpath(frozen.part("word/numbering.xml").read_tree, "./w:abstractNum | ./w:num")
        assert not any(
            xpath(p.read_tree, ".//w:numPr") for p in frozen.parts() if p.path.endswith(".xml")
        )
    else:
        reference = (
            disabled.format_bindings.heading[1]
            if format_name == "heading"
            else disabled.format_bindings.body
            if format_name == "body"
            else disabled.format_bindings.tables[0]
        )
        prototype = next(p for p in disabled.prototypes if p.prototype_id == reference)
        tree = parse((tmp_path / "disabled" / prototype.path).read_bytes())
        assert all(
            (e.get(qn("w:val")) or "").startswith("DFNeutral")
            for e in xpath(tree, ".//w:pStyle | .//w:rStyle | .//w:tblStyle")
        )
        if format_name == "body":
            assert all(e.get(qn("w:val")) == "0" for e in xpath(tree, ".//w:b"))


def test_review_remaining_section_inherits_first_even_headers(tmp_path: Path) -> None:
    from docx.enum.section import WD_SECTION_START

    source = tmp_path / "source.docx"
    document = Document()
    document.settings.odd_and_even_pages_header_footer = True
    first = document.sections[0]
    first.different_first_page_header_footer = True
    first.header.paragraphs[0].text = "OLD_DEFAULT_HEADER"
    first.first_page_header.paragraphs[0].text = "OLD_FIRST_HEADER"
    first.even_page_header.paragraphs[0].text = "OLD_EVEN_HEADER"
    document.add_paragraph("OLD_COVER", "Title")
    second = document.add_section(WD_SECTION_START.NEW_PAGE)
    second.different_first_page_header_footer = True
    document.add_heading("OLD_HEADING", 1)
    document.add_paragraph("OLD_BODY")
    document.save(source)
    selection = default_selection(inspect_template(source, import_id="review"))
    selection.ranges = [r for r in selection.ranges if r.kind != "cover"]
    freeze_template(source, selection, tmp_path / "bundle", template_id="t", version=1, name="T")
    package = DocxPackage.open(tmp_path / "bundle/effective.docx")
    sections = xpath(package.part("word/document.xml").read_tree, ".//w:sectPr")
    assert len(sections) == 1
    references = xpath(sections[0], "./w:headerReference")
    assert {e.get(qn("w:type")) for e in references} == {"default", "first", "even"}
    assert xpath(sections[0], "./w:titlePg")
    assert xpath(package.part("word/settings.xml").read_tree, "./w:evenAndOddHeaders")
    assert all(r.part is not None for r in package.resolve_rels("word/document.xml", role="header"))


@pytest.mark.parametrize(
    "instruction",
    ["TOC \\b OLD_BOOKMARK", 'TOC \\t "OLD_BUSINESS,1"', 'TOC \\o "1-3" INCLUDETEXT https://old'],
)
def test_review_unsafe_toc_switches_are_rejected(tmp_path: Path, instruction: str) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source, read_only=False)
    body = xpath(package.part("word/document.xml").tree, "./w:body")[0]
    paragraph = new_tree("w:p")
    sub(paragraph, "w:fldSimple", {"w:instr": instruction})
    body.insert(2, paragraph)
    package.save(source)
    with pytest.raises(EngineError, match="E_SCHEMA_INVALID"):
        freeze(source, tmp_path / "bundle")


def test_review_explicit_neutral_formats_exist_without_source_prototypes(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("OLD_ONLY_BODY")
    document.save(source)
    version, _ = freeze(
        source, tmp_path / "bundle", formats={"heading": False, "body": False, "numbering": False}
    )
    assert set(version.format_bindings.heading) == set(range(1, 10))
    assert version.format_bindings.body is not None
    assert version.format_bindings.list_paragraph is not None


def test_review_toc_cached_run_format_is_retained(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source, read_only=False)
    body = xpath(package.part("word/document.xml").tree, "./w:body")[0]
    paragraph = new_tree("w:p")
    field = sub(paragraph, "w:fldSimple", {"w:instr": 'TOC \\o "1-3"'})
    run = sub(field, "w:r")
    sub(sub(run, "w:rPr"), "w:color", {"w:val": "123456"})
    sub(run, "w:t").text = "OLD_TOC_CACHE"
    body.insert(2, paragraph)
    package.save(source)
    version, package = freeze(source, tmp_path / "bundle")
    marker = version.format_bindings.toc[0].marker
    wrapper = xpath(
        package.part("word/document.xml").read_tree, f'.//w:sdt[w:sdtPr/w:tag[@w:val="{marker}"]]'
    )[0]
    assert xpath(wrapper, './/w:rPr/w:color[@w:val="123456"]')


def test_review_neutral_body_does_not_inherit_source_default_style(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.styles["Normal"].font.bold = True
    document.add_paragraph("OLD_DEFAULT_BOLD")
    document.save(source)
    version, package = freeze(source, tmp_path / "bundle", formats={"body": False})
    prototype = next(
        p for p in version.prototypes if p.prototype_id == version.format_bindings.body
    )
    paragraph = parse((tmp_path / "bundle" / prototype.path).read_bytes())
    reference = xpath(paragraph, "./w:pPr/w:pStyle")[0].get(qn("w:val"))
    assert reference == "DFNeutralBody"
    style = xpath(
        package.part("word/styles.xml").read_tree, f'./w:style[@w:styleId="{reference}"]'
    )[0]
    assert not xpath(style, "./w:basedOn")
    assert xpath(paragraph, './w:r/w:rPr/w:b[@w:val="0"]')


def test_review_numbering_tab_positions_are_retained(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source)
    original = [
        e.get(qn("w:pos")) for e in xpath(package.part("word/numbering.xml").read_tree, ".//w:tab")
    ]
    assert original and all(original)
    _, frozen = freeze(source, tmp_path / "bundle")
    assert [
        e.get(qn("w:pos")) for e in xpath(frozen.part("word/numbering.xml").read_tree, ".//w:tab")
    ] == original


def test_review_explicit_neutral_table_exists_without_source_table(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("OLD_BODY")
    document.save(source)
    version, _ = freeze(source, tmp_path / "bundle", formats={"tables": False})
    assert version.format_bindings.tables
    prototype = next(
        p for p in version.prototypes if p.prototype_id == version.format_bindings.tables[0]
    )
    assert prototype.source_node_id is None and prototype.columns is None
    assert prototype.column_policy == "flexible"
    assert prototype.row_path is not None
    row = parse((tmp_path / "bundle" / prototype.row_path).read_bytes())
    assert len(xpath(row, "./w:tc")) == 1
    table = parse((tmp_path / "bundle" / prototype.path).read_bytes())
    assert xpath(table, './w:tblPr/w:tblW[@w:type="dxa"]')

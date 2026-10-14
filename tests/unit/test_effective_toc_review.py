"""复杂目录的嵌套域与完整候选支持性回归。"""

from pathlib import Path

import pytest
from docx import Document

from docfidelity.contracts.errors import EngineError
from docfidelity.ooxml.doc import Element, new_tree, qn, sub, xpath
from docfidelity.ooxml.package import DocxPackage
from docfidelity.registry.candidates import inspect_template
from docfidelity.registry.freeze import freeze_template
from docfidelity.registry.selection import default_selection

pytestmark = pytest.mark.unit
INSTRUCTION = 'TOC \\o "1-3" \\h \\z \\u'


def field_char(paragraph: Element, value: str) -> None:
    sub(sub(paragraph, "w:r"), "w:fldChar", {"w:fldCharType": value})


def build_toc(source: Path, *, unsupported_middle: bool) -> None:
    document = Document()
    document.add_heading("OLD_HEADING", 1)
    document.add_paragraph("OLD_BODY")
    document.save(source)
    package = DocxPackage.open(source, read_only=False)
    body = xpath(package.part("word/document.xml").tree, "./w:body")[0]
    first, middle, last = (new_tree("w:p") for _ in range(3))
    field_char(first, "begin")
    sub(sub(first, "w:r"), "w:instrText").text = INSTRUCTION[:9]
    sub(sub(first, "w:r"), "w:instrText").text = INSTRUCTION[9:]
    field_char(first, "separate")
    sub(sub(first, "w:r"), "w:t").text = "OLD_TOC_ENTRY"
    if unsupported_middle:
        sub(sub(middle, "w:r"), "w:footnoteReference", {"w:id": "1"})
        notes = new_tree("w:footnotes")
        sub(
            sub(sub(sub(notes, "w:footnote", {"w:id": "1"}), "w:p"), "w:r"), "w:t"
        ).text = "OLD_FOOTNOTE"
        package.add_xml_part(
            "word/footnotes.xml",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
            notes,
        )
        package.attach_relationship(
            "word/document.xml",
            type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
            target="footnotes.xml",
        )
    else:
        for paragraph in (first, middle):
            field_char(paragraph, "begin")
            sub(sub(paragraph, "w:r"), "w:instrText").text = "PAGEREF _Toc123 \\h"
            field_char(paragraph, "separate")
            sub(sub(paragraph, "w:r"), "w:t").text = "OLD_PAGE_CACHE"
            field_char(paragraph, "end")
    field_char(last, "end")
    for index, paragraph in enumerate((first, middle, last)):
        body.insert(index, paragraph)
    package.save(source)


def test_nested_pageref_cache_does_not_corrupt_outer_toc_instruction(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    build_toc(source, unsupported_middle=False)
    selection = default_selection(inspect_template(source, import_id="nested"))
    assert any(r.kind == "toc" for r in selection.ranges)
    version = freeze_template(
        source, selection, tmp_path / "bundle", template_id="t", version=1, name="T"
    )
    assert version.format_bindings.toc[0].instruction == INSTRUCTION
    package = DocxPackage.open(tmp_path / "bundle/effective.docx")
    root = package.part("word/document.xml").read_tree
    assert xpath(root, './/w:fldSimple[@w:dirty="true"]')[0].get(qn("w:instr")) == INSTRUCTION
    assert not xpath(root, ".//w:instrText | .//w:fldChar")
    assert all(
        bad not in package.part("word/document.xml").to_bytes()
        for bad in (b"OLD_", b"PAGEREF", b"_Toc123")
    )


def test_unsupported_middle_skips_whole_default_toc_and_body_range(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    build_toc(source, unsupported_middle=True)
    draft = inspect_template(source, import_id="unsupported")
    candidate = next(c for c in draft.candidates if c.kind == "toc")
    assert candidate.status == "detected"
    selection = default_selection(draft)
    assert not any(r.kind == "toc" for r in selection.ranges)
    nodes = {n.node_id: n for n in draft.nodes}
    assert all(nodes[r.start_node_id].block_index > 2 for r in selection.ranges)
    version = freeze_template(
        source, selection, tmp_path / "bundle", template_id="t", version=1, name="T"
    )
    assert not version.format_bindings.toc
    selection.ranges.extend(candidate.ranges)
    with pytest.raises(EngineError, match="E_SCHEMA_INVALID"):
        freeze_template(
            source, selection, tmp_path / "explicit", template_id="t", version=2, name="T"
        )

"""材料图片尺寸、模型隔离和非正文对象的跨层回归。"""

import binascii
import json
import struct
import zlib
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from pydantic import ValidationError

from docfidelity.compose.assembly import assemble_content
from docfidelity.compose.draft import generate_section
from docfidelity.compose.outline import generate_outline, suggest_outline
from docfidelity.contracts.documents import DocumentDraft, DraftBlock, SectionDraft
from docfidelity.contracts.generation import WritingRequest, WritingResult
from docfidelity.contracts.ir import FigureData
from docfidelity.ooxml.doc import qn, xpath
from docfidelity.ooxml.package import DocxPackage
from docfidelity.pipeline.materials import ingest_material
from docfidelity.render.effective import render_effective
from tests.unit.test_content_adapter_v2 import generator, response
from tests.unit.test_render_effective import _bundle

pytestmark = pytest.mark.unit


def png_bytes():
    def chunk(kind, data):
        return (
            struct.pack(">I", len(data))
            + kind
            + data
            + struct.pack(">I", binascii.crc32(kind + data))
        )

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 800, 400, 8, 0, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress((b"\0" + b"\x80" * 800) * 400))
        + chunk(b"IEND", b"")
    )


def image_docx(*, extent="valid", suffix="png", nonbody=None):
    document = Document()
    document.add_heading("图片章节", 1)
    document.add_paragraph("保留正文。")
    parent = document.sections[0].header if nonbody == "header" else document
    for width in [914400, 1828800]:
        parent.add_paragraph().add_run().add_picture(BytesIO(png_bytes()), width=width)
    if nonbody == "footer":
        document.sections[0].footer.paragraphs[0]._p.append(OxmlElement("w:object"))
    stream = BytesIO()
    document.save(stream)
    output = BytesIO()
    with ZipFile(stream) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "word/document.xml":
                root = etree.fromstring(data)
                for node in root.iter(qn("wp:extent")):
                    if extent == "missing":
                        node.getparent().remove(node)
                    elif extent != "valid":
                        node.set("cx", extent)
                data = etree.tostring(root)
            if suffix != "png":
                if name.endswith("image1.png"):
                    name = name[:-3] + suffix
                    if suffix == "bmp":
                        data = b"BM" + b"\0" * 52 + b"\xff" * 4
                    elif suffix == "tiff":
                        data = b"II*\0\x08\0\0\0\0\0\0\0\0\0"
                    else:
                        data = b"RIFF\x04\0\0\0WEBP"
                elif name.endswith(".rels"):
                    data = data.replace(b"image1.png", ("image1." + suffix).encode())
            target.writestr(name, data)
    return output.getvalue()


def draft(material, **updates):
    return DocumentDraft.model_validate(
        dict(
            document_id="d",
            revision=1,
            name="新文档",
            template_id="t",
            template_version=1,
            material_ids=[material.material_id],
            mode="format_only",
            outline=suggest_outline([material]),
            outline_confirmed=True,
        )
        | updates
    )


@pytest.mark.parametrize("mode", ["format_only", "organize"])
def test_displayed_sizes_follow_source_occurrences_through_actual_render(tmp_path, mode):
    source, assets = ingest_material(image_docx(), name="new.docx", material_id="m")
    original = source.model_dump_json()
    figures = [b for b in source.blocks if b.kind == "figure"]
    assert figures[0].image_ref == figures[1].image_ref
    doc = draft(source, mode=mode)
    section = generate_section(doc, [source], doc.outline[0].section_id, None).section
    doc.sections = [section]
    content = assemble_content(doc, [source])
    directory, _ = _bundle(tmp_path)
    for key, data in assets.items():
        path = tmp_path / "assets" / key
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
        assert data == png_bytes()
    output = render_effective(directory, content, {}, tmp_path / "new.docx")
    package = DocxPackage.open(output)
    sizes = [
        (int(n.get("cx")), int(n.get("cy")))
        for n in xpath(package.part("word/document.xml").read_tree, ".//wp:extent")
    ]
    assert sizes == [(914400, 457200), (1828800, 914400)]
    assert source.model_dump_json() == original
    assert all(
        package.part(name).to_bytes() == png_bytes()
        for name in package.part_names()
        if name.startswith("word/media/")
    )


def test_actual_model_payload_and_draft_cannot_receive_or_set_geometry():
    source, _ = ingest_material(image_docx(), name="new.docx", material_id="m")
    doc = draft(source, mode="organize")
    calls = []

    def handler(req):
        calls.append(json.loads(req.content))
        return response(json.dumps({"outline": [s.model_dump() for s in doc.outline]}))

    result = generate_outline(doc, [source], generator(handler))
    assert not result.error
    assert (
        next(b for b in source.blocks if b.kind == "figure").model_dump().get("image_extent")
        is not None
    )
    user = calls[0]["messages"][1]["content"]
    assert "image_extent" not in user and "914400" not in user
    assert "image_extent" not in json.dumps(WritingRequest.model_json_schema())
    figure = next(b for b in source.blocks if b.kind == "figure")
    with pytest.raises(ValidationError):
        DraftBlock(kind="figure", image_ref=figure.image_ref, image_extent={"cx": 1, "cy": 2})


@pytest.mark.parametrize("extent", ["missing", "0", "-1", "broken"])
@pytest.mark.parametrize("suffix", ["png", "bmp", "tiff", "webp"])
def test_unusable_image_is_pending_at_ingest_without_late_build_failure(extent, suffix):
    source, assets = ingest_material(
        image_docx(extent=extent, suffix=suffix), name="new.docx", material_id="m"
    )
    assert not any(b.kind == "figure" for b in source.blocks)
    assert not assets
    assert any("word/document.xml/body/3" in p and "待确认" in p for p in source.pending)
    assert assemble_content(draft(source), [source]).unresolved


@pytest.mark.parametrize("suffix", ["bmp", "tiff"])
def test_renderable_formats_use_trusted_display_extent(tmp_path, suffix):
    source, assets = ingest_material(image_docx(suffix=suffix), name="new.docx", material_id="m")
    assert len([b for b in source.blocks if b.kind == "figure"]) == 2
    content = assemble_content(draft(source), [source])
    directory, _ = _bundle(tmp_path)
    for key, data in assets.items():
        path = tmp_path / "assets" / key
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
    output = render_effective(directory, content, {}, tmp_path / "new.docx")
    assert [(s.width, s.height) for s in Document(output).inline_shapes] == [
        (914400, 457200),
        (1828800, 914400),
    ]


def test_webp_is_pending_even_with_trusted_display_extent():
    source, assets = ingest_material(image_docx(suffix="webp"), name="new.docx", material_id="m")
    assert not assets and not any(b.kind == "figure" for b in source.blocks)
    assert any(".webp" in p and "body/3" in p for p in source.pending)


@pytest.mark.parametrize("nonbody,kind", [("header", "drawing"), ("footer", "object")])
def test_nonbody_objects_without_text_are_locatable_pending(nonbody, kind):
    source, _ = ingest_material(image_docx(nonbody=nonbody), name="new.docx", material_id="m")
    assert any("word/" + nonbody in p and kind in p and "待确认" in p for p in source.pending)
    assert any(b.text == "保留正文。" for b in source.blocks)


@pytest.mark.parametrize(
    "fields",
    [
        {"image_extent": {"cx": 1, "cy": 0}},
        {"image_extent": {"cx": 1}},
        {"image_extent": {"cx": 1, "cy": 2}, "width_px": 3, "height_px": 4},
    ],
)
def test_ir_rejects_invalid_or_ambiguous_display_geometry(fields):
    with pytest.raises(ValidationError):
        FigureData(image_ref="images/a.png", **fields)


def test_manual_unsourced_picture_cannot_enlarge_trusted_asset():
    source, _ = ingest_material(image_docx(), name="new.docx", material_id="m")
    doc = draft(source)
    ref = next(b.image_ref for b in source.blocks if b.kind == "figure")
    doc.sections = [
        SectionDraft(
            section_id=doc.outline[0].section_id,
            edited=True,
            blocks=[DraftBlock(kind="figure", image_ref=ref)],
        )
    ]
    content = assemble_content(doc, [source])
    assert content.units[0].children[0].fields["image_extent"] == {"cx": 914400, "cy": 457200}


def test_manual_cited_picture_keeps_its_own_occurrence_size():
    source, _ = ingest_material(image_docx(), name="new.docx", material_id="m")
    doc = draft(source)
    figure = [b for b in source.blocks if b.kind == "figure"][-1]
    doc.sections = [
        SectionDraft(
            section_id=doc.outline[0].section_id,
            edited=True,
            blocks=[
                DraftBlock(
                    kind="figure", image_ref=figure.image_ref, source_refs=[f"m/{figure.block_id}"]
                )
            ],
        )
    ]
    assert assemble_content(doc, [source]).units[0].children[0].fields["image_extent"] == {
        "cx": 1828800,
        "cy": 914400,
    }


def test_organized_figures_keep_cited_geometry_outside_model_draft():
    source, _ = ingest_material(image_docx(), name="new.docx", material_id="m")
    doc = draft(source, mode="organize")
    section = generate_section(doc, [source], doc.outline[0].section_id, None).section

    class Provider:
        def generate(self, request):
            assert all(not hasattr(b, "image_extent") for m in request.materials for b in m.blocks)
            return WritingResult(section=section)

    result = generate_section(doc, [source], doc.outline[0].section_id, Provider())
    assert not result.error
    doc.sections = [result.section]
    figures = [u for u in assemble_content(doc, [source]).walk() if u.type == "figure"]
    assert [u.fields["image_extent"] for u in figures] == [
        {"cx": 914400, "cy": 457200},
        {"cx": 1828800, "cy": 914400},
    ]


def test_old_material_snapshots_without_geometry_remain_readable():
    from docfidelity.contracts.materials import SourceMaterial

    source, _ = ingest_material(image_docx(), name="new.docx", material_id="m")
    old = SourceMaterial.model_validate(
        source.model_dump(exclude={"blocks": {"__all__": {"image_extent"}}})
    )
    content = assemble_content(draft(old), [old])
    assert [u.type for u in content.walk()].count("figure") == 2


def test_provider_geometry_output_is_rejected_instead_of_adopted():
    source, _ = ingest_material(image_docx(), name="new.docx", material_id="m")
    doc = draft(source, mode="organize")
    section = generate_section(doc, [source], doc.outline[0].section_id, None).section
    payload = section.model_dump()
    next(b for b in payload["blocks"] if b["kind"] == "figure")["image_extent"] = {
        "cx": 99999999,
        "cy": 99999999,
    }
    result = generate_section(
        doc,
        [source],
        doc.outline[0].section_id,
        generator(lambda req: response(json.dumps({"section": payload}))),
    )
    assert result.error and result.section == section

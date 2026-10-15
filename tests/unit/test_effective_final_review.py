"""最终审查：自制 DOCX 的冻结与渲染继承语义回归。"""

from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE

from docfidelity.contracts.imports import FormatSelection, VariableMapping
from docfidelity.contracts.ir import ContentIR, ContentUnit
from docfidelity.ooxml.doc import Element, new_tree, parse, qn, sub, xpath
from docfidelity.ooxml.package import DocxPackage
from docfidelity.registry.candidates import inspect_template
from docfidelity.registry.freeze import freeze_template
from docfidelity.registry.selection import default_selection
from docfidelity.render.effective import render_effective

pytestmark = pytest.mark.unit

_ROLES = (
    ("header", "default", "header"),
    ("header", "first", "first_page_header"),
    ("header", "even", "even_page_header"),
    ("footer", "default", "footer"),
    ("footer", "first", "first_page_footer"),
    ("footer", "even", "even_page_footer"),
)


def _ir(*units: ContentUnit) -> ContentIR:
    return ContentIR(project_id="p", doc_key="d", template_ref="t@1", units=list(units))


def _section_roles(package: DocxPackage) -> list[dict[tuple[str, str], str]]:
    relationships = {r.rel_id: r for r in package.relationships("word/document.xml")}
    inherited: dict[tuple[str, str], str] = {}
    result = []
    for section in xpath(package.part("word/document.xml").read_tree, ".//w:sectPr"):
        for kind, role, _ in _ROLES:
            for reference in xpath(section, f'./w:{kind}Reference[@w:type="{role}"]'):
                relation = relationships[reference.get(qn("r:id"))]
                path = "word/" + relation.target
                assert package.has_part(path)
                inherited[kind, role] = "".join(
                    e.text or "" for e in xpath(package.part(path).read_tree, ".//w:t")
                )
        result.append(dict(inherited))
    return result


@pytest.mark.parametrize("remove_middle", [False, True])
@pytest.mark.parametrize("cancel_middle", [False, True])
@pytest.mark.parametrize("page_setup", [False, True])
def test_cancelled_six_section_roles_do_not_reinherit_previous_content(
    tmp_path: Path, remove_middle: bool, cancel_middle: bool, page_setup: bool
) -> None:
    document = Document()
    document.settings.odd_and_even_pages_header_footer = True
    for index in range(3):
        section = (
            document.sections[0] if index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
        )
        section.different_first_page_header_footer = True
        document.add_paragraph(f"SECTION_{index}")
        if index < 2:
            for kind, role, attribute in _ROLES:
                part = getattr(section, attribute)
                part.is_linked_to_previous = False
                part.paragraphs[0].text = f"ROLE_{index}_{kind}_{role}"
    source = tmp_path / "source.docx"
    document.save(source)
    draft = inspect_template(source, import_id="review")
    selection = default_selection(draft)
    selection.formats.page_setup = page_setup
    by_id = {node.node_id: node for node in draft.nodes}
    source_body = xpath(DocxPackage.open(source).part("word/document.xml").read_tree, "./w:body")[0]
    boundaries = [
        i
        for i, block in enumerate(source_body)
        if block.tag == qn("w:sectPr") or xpath(block, "./w:pPr/w:sectPr")
    ]
    for span in list(selection.ranges):
        node = by_id[span.start_node_id]
        if (cancel_middle and node.text_preview.startswith("ROLE_1_")) or (
            remove_middle
            and node.part == "word/document.xml"
            and boundaries[0] < node.block_index <= boundaries[1]
        ):
            selection.ranges.remove(span)
    values = {}
    for node in draft.nodes:
        if node.kind == "paragraph" and node.text_preview.startswith("ROLE_"):
            if cancel_middle and node.text_preview.startswith("ROLE_1_"):
                continue
            key = node.text_preview
            selection.variables.append(VariableMapping(key=key, node_id=node.node_id))
            values[key] = key
    directory = tmp_path / "bundle"
    freeze_template(source, selection, directory, template_id="t", version=1, name="T")
    source.unlink()
    output = render_effective(directory, _ir(), values, tmp_path / "out.docx")
    for path in (directory / "effective.docx", output):
        package = DocxPackage.open(path)
        roles = _section_roles(package)
        assert len(roles) == (2 if remove_middle else 3)
        for kind, role, _ in _ROLES:
            expected = "" if cancel_middle else f"ROLE_1_{kind}_{role}"
            # 冻结时变量仍为空，渲染后实际保留文字能暴露继承错误。
            if path == output:
                assert roles[0][kind, role] == f"ROLE_0_{kind}_{role}"
                assert roles[-1][kind, role] == expected
                if not remove_middle:
                    assert roles[1][kind, role] == expected
            if cancel_middle:
                sections = xpath(package.part("word/document.xml").read_tree, ".//w:sectPr")
                assert xpath(sections[1], f'./w:{kind}Reference[@w:type="{role}"]')
        assert xpath(package.part("word/settings.xml").read_tree, "./w:evenAndOddHeaders")
        assert all(
            xpath(s, "./w:titlePg")
            for s in xpath(package.part("word/document.xml").read_tree, ".//w:sectPr")
        )


@pytest.mark.parametrize("body_kind", ["ordinary", "name_only", "cancel_style", "cancel_direct"])
@pytest.mark.parametrize("list_first", [False, True])
def test_list_binding_uses_effective_numbering_and_body_never_uses_list(
    tmp_path: Path, body_kind: str, list_first: bool
) -> None:
    document = Document()
    parent = document.styles.add_style("BulletParent", WD_STYLE_TYPE.PARAGRAPH)
    parent.base_style = document.styles["List Bullet"]
    custom = document.styles.add_style("CustomBullet", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = parent
    body_style = document.styles.add_style("ListOrdinary", WD_STYLE_TYPE.PARAGRAPH)
    body_style.base_style = document.styles["Normal"]
    cancelled = document.styles.add_style("CancelledBullet", WD_STYLE_TYPE.PARAGRAPH)
    cancelled.base_style = custom
    sub(sub(sub(cancelled.element, "w:pPr"), "w:numPr"), "w:numId", {"w:val": "0"})

    def add_body() -> None:
        style = {
            "ordinary": "Normal",
            "name_only": "ListOrdinary",
            "cancel_style": "CancelledBullet",
            "cancel_direct": "CustomBullet",
        }[body_kind]
        paragraph = document.add_paragraph("BODY", style)
        if body_kind == "cancel_direct":
            sub(sub(paragraph._p.get_or_add_pPr(), "w:numPr"), "w:numId", {"w:val": "0"})

    if not list_first:
        add_body()
    document.add_paragraph("BULLET", custom)
    if list_first:
        add_body()
    source = tmp_path / "source.docx"
    document.save(source)
    draft = inspect_template(source, import_id="review")
    directory = tmp_path / "bundle"
    version = freeze_template(
        source, default_selection(draft), directory, template_id="t", version=1, name="T"
    )
    nodes = {node.node_id: node for node in draft.nodes}
    prototypes = {p.prototype_id: p for p in version.prototypes}
    assert version.format_bindings.list_paragraph is not None
    assert (
        nodes[prototypes[version.format_bindings.list_paragraph].source_node_id].text_preview
        == "BULLET"
    )
    assert nodes[prototypes[version.format_bindings.body].source_node_id].text_preview == "BODY"
    source.unlink()
    output = render_effective(
        directory,
        _ir(
            ContentUnit(unit_id="p", type="para", text="Fresh body"),
            ContentUnit(unit_id="l", type="bullets", text="Fresh bullet"),
        ),
        {},
        tmp_path / "out.docx",
    )
    package = DocxPackage.open(output)
    root = package.part("word/document.xml").read_tree
    bullet = xpath(root, './/w:p[w:r/w:t="Fresh bullet"]')[0]
    assert xpath(bullet, "./w:pPr/w:numPr/w:numId")
    body = xpath(root, './/w:p[w:r/w:t="Fresh body"]')[0]
    assert (
        not xpath(body, './w:pPr/w:pStyle[@w:val="CustomBullet"]') or body_kind == "cancel_direct"
    )
    num = _effective_properties(package, body, "pPr").get("numPr", {})
    assert num.get("numId", "0") == "0"


@pytest.mark.parametrize("default_location", ["style", "document"])
def test_implicit_default_numbering_survives_freeze_and_render(
    tmp_path: Path, default_location: str
) -> None:
    document = Document()
    document.add_paragraph("BULLET")
    body = document.add_paragraph("BODY")
    sub(sub(body._p.get_or_add_pPr(), "w:numPr"), "w:numId", {"w:val": "0"})
    source = tmp_path / "source.docx"
    document.save(source)
    package = DocxPackage.open(source, read_only=False)
    styles = package.part("word/styles.xml").tree
    num_id = xpath(styles, './w:style[@w:styleId="ListBullet"]/w:pPr/w:numPr/w:numId')[0].get(
        qn("w:val")
    )
    if default_location == "style":
        parent = xpath(styles, './w:style[@w:styleId="Normal"]')[0]
        ppr = sub(parent, "w:pPr")
    else:
        ppr = xpath(styles, "./w:docDefaults/w:pPrDefault/w:pPr")[0]
    sub(sub(ppr, "w:numPr"), "w:numId", {"w:val": num_id})
    package.save(source)
    selection = default_selection(inspect_template(source, import_id="review"))
    directory = tmp_path / "bundle"
    version = freeze_template(source, selection, directory, template_id="t", version=1, name="T")
    assert version.format_bindings.list_paragraph is not None
    source.unlink()
    output = render_effective(
        directory,
        _ir(ContentUnit(unit_id="l", type="bullets", text="Fresh bullet")),
        {},
        tmp_path / "out.docx",
    )
    assert xpath(
        DocxPackage.open(output).part("word/document.xml").read_tree,
        './/w:p[w:r/w:t="Fresh bullet"]/w:pPr/w:numPr/w:numId',
    )


def test_numbering_disabled_does_not_restore_document_default_pointer(tmp_path: Path) -> None:
    document = Document()
    body = document.add_paragraph("BODY")
    sub(sub(body._p.get_or_add_pPr(), "w:numPr"), "w:numId", {"w:val": "0"})
    document.add_paragraph("BULLET")
    source = tmp_path / "source.docx"
    document.save(source)
    package = DocxPackage.open(source, read_only=False)
    styles = package.part("word/styles.xml").tree
    pointer = xpath(styles, './w:style[@w:styleId="ListBullet"]/w:pPr/w:numPr/w:numId')[0]
    defaults = xpath(styles, "./w:docDefaults/w:pPrDefault/w:pPr")[0]
    sub(sub(defaults, "w:numPr"), "w:numId", {"w:val": pointer.get(qn("w:val"))})
    package.save(source)
    selection = default_selection(inspect_template(source, import_id="review"))
    selection.formats.numbering = False
    directory = tmp_path / "bundle"
    version = freeze_template(source, selection, directory, template_id="t", version=1, name="T")
    frozen = DocxPackage.open(directory / "effective.docx")
    for part in frozen.parts():
        if part.path.endswith(".xml"):
            assert not xpath(part.read_tree, ".//w:numPr"), part.path
    for prototype in version.prototypes:
        assert not xpath(parse((directory / prototype.path).read_bytes()), ".//w:numPr")
    output = render_effective(
        directory,
        _ir(
            ContentUnit(unit_id="p", type="para", text="Fresh body"),
            ContentUnit(unit_id="l", type="numbered", text="Fresh item"),
        ),
        {},
        tmp_path / "out.docx",
    )
    root = DocxPackage.open(output).part("word/document.xml").read_tree
    assert not xpath(root, './/w:p[w:r/w:t="Fresh body"]/w:pPr/w:numPr')
    assert xpath(root, './/w:p[w:r/w:t="Fresh item"]/w:pPr/w:numPr')


def _effective_properties(
    package: DocxPackage, paragraph: Element, kind: str
) -> dict[str, dict[str, str]]:
    """仅评估夹具涉及的复合属性；按实际包的默认值、样式链、直接格式叠加。"""
    styles = package.part("word/styles.xml").read_tree
    result: dict[str, dict[str, str]] = {}

    def merge(properties: Element, *, style: bool = False) -> None:
        for child in properties:
            tag = child.tag.rpartition("}")[2]
            values = {key.rpartition("}")[2]: value for key, value in child.attrib.items()}
            if not values:
                values = {"val": "1"}
            if style and tag == "vanish":
                if values["val"] in {"0", "false", "off"}:
                    continue
                values["val"] = "0" if result.get(tag, {}).get("val", "0") == "1" else "1"
            if tag == "numPr":
                values = {e.tag.rpartition("}")[2]: e.get(qn("w:val"), "") for e in child}
            result.setdefault(tag, {}).update(values)

    for defaults in xpath(styles, f"./w:docDefaults/w:{kind}Default/w:{kind}"):
        merge(defaults)
    tables = list(paragraph.iterancestors(qn("w:tbl")))
    if tables:
        table_identifiers = [
            reference.get(qn("w:val")) for reference in xpath(tables[0], "./w:tblPr/w:tblStyle")
        ]
        if not table_identifiers:
            table_identifiers = [
                style.get(qn("w:styleId"))
                for style in xpath(styles, './w:style[@w:type="table"]')
                if style.get(qn("w:default")) in {"1", "true", "on"}
            ]
        for table_identifier in table_identifiers:
            table_chain = []
            while table_identifier:
                current = xpath(styles, f'./w:style[@w:styleId="{table_identifier}"]')[0]
                table_chain.insert(0, current)
                base = xpath(current, "./w:basedOn")
                table_identifier = base[0].get(qn("w:val")) if base else None
            for style in table_chain:
                for properties in xpath(style, f"./w:{kind}"):
                    merge(properties, style=True)
            row = next(paragraph.iterancestors(qn("w:tr")), None)
            if row is not None and row is xpath(tables[0], "./w:tr")[0]:
                for style in table_chain:
                    for properties in xpath(style, f'./w:tblStylePr[@w:type="firstRow"]/w:{kind}'):
                        merge(properties, style=True)
    for reference in xpath(paragraph, "./w:pPr/w:numPr/w:numId"):
        numbering = package.part("word/numbering.xml").read_tree
        identifier = reference.get(qn("w:val"))
        for pointer in xpath(numbering, f'./w:num[@w:numId="{identifier}"]/w:abstractNumId'):
            for properties in xpath(
                numbering,
                f'./w:abstractNum[@w:abstractNumId="{pointer.get(qn("w:val"))}"]/w:lvl[@w:ilvl="0"]/w:{kind}',
            ):
                merge(properties, style=True)
    references = xpath(paragraph, "./w:pPr/w:pStyle")
    identifiers = [e.get(qn("w:val")) for e in references]
    if not identifiers:
        identifiers = [
            e.get(qn("w:styleId"))
            for e in xpath(styles, './w:style[@w:type="paragraph"][@w:default="1"]')
        ]
    chain = []
    while identifiers:
        identifier = identifiers.pop()
        current = xpath(styles, f'./w:style[@w:styleId="{identifier}"]')
        if not current:
            break
        chain.append(current[0])
        identifiers = [e.get(qn("w:val")) for e in xpath(current[0], "./w:basedOn")]
    for style in reversed(chain):
        for properties in xpath(style, f"./w:{kind}"):
            merge(properties, style=True)
    for properties in xpath(paragraph, "./w:pPr" if kind == "pPr" else "./w:r/w:rPr"):
        merge(properties)
    return result


@pytest.mark.parametrize(
    "body,heading", [(True, True), (False, True), (True, False), (False, False)]
)
@pytest.mark.parametrize("ancestor_override", [False, True])
def test_neutral_formats_isolate_defaults_and_keep_enabled_style_chain(
    tmp_path: Path, body: bool, heading: bool, ancestor_override: bool
) -> None:
    document = Document()
    ancestor = document.styles.add_style("SourceAncestor", WD_STYLE_TYPE.PARAGRAPH)
    ancestor.base_style = document.styles["Normal"]
    sub(sub(ancestor.element, "w:pPr"), "w:ind", {"w:firstLine": "180"})
    run_properties = sub(ancestor.element, "w:rPr")
    sub(run_properties, "w:rFonts", {"w:ascii": "AncestorFont"})
    if ancestor_override:
        sub(run_properties, "w:vanish")
    document.styles["Heading 1"].base_style = ancestor
    heading_properties = document.styles["Heading 1"].element.find(qn("w:rPr"))
    if heading_properties is not None:
        heading_fonts = heading_properties.find(qn("w:rFonts"))
        if heading_fonts is not None:
            heading_properties.remove(heading_fonts)
    child = document.styles.add_style("SourceBody", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = ancestor
    document.add_heading("HEADING", 1)
    document.add_paragraph("BODY", child)
    source = tmp_path / "source.docx"
    document.save(source)
    package = DocxPackage.open(source, read_only=False)
    styles = package.part("word/styles.xml").tree
    for old in xpath(styles, "./w:docDefaults"):
        styles.remove(old)
    defaults = new_tree("w:docDefaults")
    styles.insert(0, defaults)
    rp = sub(sub(defaults, "w:rPrDefault"), "w:rPr")
    sub(rp, "w:vanish")
    sub(
        rp,
        "w:rFonts",
        {"w:ascii": "DefaultAscii", "w:hAnsi": "DefaultLatin", "w:eastAsia": "DefaultEast"},
    )
    pp = sub(sub(defaults, "w:pPrDefault"), "w:pPr")
    sub(pp, "w:ind", {"w:left": "720", "w:right": "360"})
    sub(sub(pp, "w:pBdr"), "w:bottom", {"w:val": "single", "w:sz": "8"})
    package.save(source)
    selection = default_selection(inspect_template(source, import_id="review"))
    selection.formats = FormatSelection(body=body, heading=heading)
    directory = tmp_path / "bundle"
    freeze_template(source, selection, directory, template_id="t", version=1, name="T")
    source.unlink()
    output = render_effective(
        directory,
        _ir(
            ContentUnit(unit_id="h", type="heading", level=1, text="Fresh heading"),
            ContentUnit(unit_id="p", type="para", text="Fresh body"),
        ),
        {},
        tmp_path / "out.docx",
    )
    package = DocxPackage.open(output)
    root = package.part("word/document.xml").read_tree
    for text, enabled in (("Fresh heading", heading), ("Fresh body", body)):
        paragraph = xpath(root, f'.//w:p[w:r/w:t="{text}"]')[0]
        rpr = _effective_properties(package, paragraph, "rPr")
        ppr = _effective_properties(package, paragraph, "pPr")
        if enabled:
            assert rpr["vanish"]["val"] == ("0" if ancestor_override else "1")
            assert rpr["rFonts"]["ascii"] == "AncestorFont"
            assert rpr["rFonts"]["hAnsi"] == "DefaultLatin"
            assert rpr["rFonts"]["eastAsia"] == "DefaultEast"
            assert ppr["ind"] == {"left": "720", "right": "360", "firstLine": "180"}
            assert "pBdr" in ppr
        else:
            assert rpr.get("vanish", {}).get("val", "0") == "0"
            assert rpr["rFonts"]["ascii"] == "Arial"
            assert "ind" not in ppr
            assert "pBdr" not in ppr


@pytest.mark.parametrize("body_enabled", [False, True])
def test_defaults_isolation_preserves_table_and_numbering_precedence(
    tmp_path: Path, body_enabled: bool
) -> None:
    document = Document()
    document.add_heading("HEADING", 1)
    document.add_paragraph("BODY")
    document.add_paragraph("BULLET", "List Bullet")
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.cell(0, 0).text = "CELL"
    source = tmp_path / "source.docx"
    document.save(source)
    package = DocxPackage.open(source, read_only=False)
    styles = package.part("word/styles.xml").tree
    defaults = xpath(styles, "./w:docDefaults")[0]
    rp = xpath(defaults, "./w:rPrDefault/w:rPr")[0]
    for size in xpath(rp, "./w:sz"):
        rp.remove(size)
    sub(rp, "w:sz", {"w:val": "48"})
    sub(rp, "w:vanish")
    pp = xpath(defaults, "./w:pPrDefault/w:pPr")[0]
    sub(pp, "w:ind", {"w:left": "720"})
    table_style = xpath(styles, './w:style[@w:styleId="TableGrid"]')[0]
    sub(sub(table_style, "w:rPr"), "w:sz", {"w:val": "28"})
    conditional = sub(table_style, "w:tblStylePr", {"w:type": "firstRow"})
    sub(sub(conditional, "w:rPr"), "w:sz", {"w:val": "32"})
    list_style = xpath(styles, './w:style[@w:styleId="ListBullet"]')[0]
    number_id = xpath(list_style, "./w:pPr/w:numPr/w:numId")[0].get(qn("w:val"))
    numbering = package.part("word/numbering.xml").tree
    abstract = xpath(numbering, f'./w:num[@w:numId="{number_id}"]/w:abstractNumId')[0].get(
        qn("w:val")
    )
    level = xpath(numbering, f'./w:abstractNum[@w:abstractNumId="{abstract}"]/w:lvl[@w:ilvl="0"]')[
        0
    ]
    indent = xpath(level, "./w:pPr/w:ind")[0]
    indent.set(qn("w:left"), "1440")
    package.save(source)
    selection = default_selection(inspect_template(source, import_id="review"))
    selection.formats.heading = False
    selection.formats.body = body_enabled
    directory = tmp_path / "bundle"
    freeze_template(source, selection, directory, template_id="t", version=1, name="T")
    output = render_effective(
        directory,
        _ir(
            ContentUnit(unit_id="p", type="para", text="Fresh body"),
            ContentUnit(unit_id="l", type="bullets", text="Fresh bullet"),
            ContentUnit(
                unit_id="t", type="table", fields={"columns": ["A"], "rows": [["Fresh cell"]]}
            ),
        ),
        {},
        tmp_path / "out.docx",
    )
    frozen = DocxPackage.open(output)
    root = frozen.part("word/document.xml").read_tree
    cell = xpath(root, './/w:p[w:r/w:t="Fresh cell"]')[0]
    header = xpath(root, './/w:p[w:r/w:t="A"]')[0]
    bullet = xpath(root, './/w:p[w:r/w:t="Fresh bullet"]')[0]
    paragraph = xpath(root, './/w:p[w:r/w:t="Fresh body"]')[0]
    assert _effective_properties(frozen, cell, "rPr")["sz"]["val"] == (
        "28" if body_enabled else "22"
    )
    assert _effective_properties(frozen, header, "rPr")["sz"]["val"] == (
        "32" if body_enabled else "22"
    )
    assert _effective_properties(frozen, paragraph, "rPr")["sz"]["val"] == (
        "48" if body_enabled else "22"
    )
    for node in (cell, header, paragraph, bullet):
        assert _effective_properties(frozen, node, "rPr").get("vanish", {}).get("val", "0") == (
            "1" if body_enabled else "0"
        )
    if body_enabled:
        assert _effective_properties(frozen, bullet, "pPr")["ind"]["left"] == "1440"
        assert _effective_properties(frozen, paragraph, "pPr")["ind"]["left"] == "720"


@pytest.mark.parametrize("tables_enabled", [False, True])
def test_neutral_table_isolates_default_table_style_after_freeze_and_render(
    tmp_path: Path, tables_enabled: bool
) -> None:
    document = Document()
    document.add_paragraph("BODY")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "CELL"
    source = tmp_path / "source.docx"
    document.save(source)
    package = DocxPackage.open(source, read_only=False)
    styles = package.part("word/styles.xml").tree
    defaults = xpath(styles, "./w:docDefaults")[0]
    sub(xpath(defaults, "./w:rPrDefault/w:rPr")[0], "w:vanish")
    ppr = xpath(defaults, "./w:pPrDefault/w:pPr")[0]
    sub(sub(ppr, "w:pBdr"), "w:bottom", {"w:val": "single", "w:sz": "8"})
    package.save(source)
    selection = default_selection(inspect_template(source, import_id="review"))
    selection.formats = FormatSelection(tables=tables_enabled, body=True, heading=False)
    directory = tmp_path / "bundle"
    version = freeze_template(source, selection, directory, template_id="t", version=1, name="T")
    frozen = DocxPackage.open(directory / "effective.docx")
    prototype = next(
        p for p in version.prototypes if p.prototype_id == version.format_bindings.tables[0]
    )
    table = parse((directory / prototype.path).read_bytes())
    for paragraph in xpath(table, ".//w:p"):
        assert _effective_properties(frozen, paragraph, "rPr").get("vanish", {}).get(
            "val", "0"
        ) == ("1" if tables_enabled else "0")
        assert ("pBdr" in _effective_properties(frozen, paragraph, "pPr")) == tables_enabled
    source.unlink()
    output = render_effective(
        directory,
        _ir(
            ContentUnit(
                unit_id="t", type="table", fields={"columns": ["A"], "rows": [["Fresh cell"]]}
            )
        ),
        {},
        tmp_path / "out.docx",
    )
    rendered = DocxPackage.open(output)
    root = rendered.part("word/document.xml").read_tree
    for text in ("A", "Fresh cell"):
        paragraph = xpath(root, f'.//w:p[w:r/w:t="{text}"]')[0]
        assert _effective_properties(rendered, paragraph, "rPr").get("vanish", {}).get(
            "val", "0"
        ) == ("1" if tables_enabled else "0")
        assert ("pBdr" in _effective_properties(rendered, paragraph, "pPr")) == tables_enabled

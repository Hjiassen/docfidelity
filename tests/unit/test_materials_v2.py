"""新版材料接入必须保留正文、定位和不支持对象的告警。"""

from __future__ import annotations

import base64
import hashlib
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement

from docfidelity.contracts.errors import EngineError
from docfidelity.ingest.materials import ingest_text
from docfidelity.pipeline.materials import ingest_material

pytestmark = pytest.mark.unit


def test_markdown_preserves_structure_unicode_and_locations() -> None:
    tick = chr(96)
    raw = (
        "# 新材料\n\n正文 **加粗** 与 "
        + tick
        + "代码"
        + tick
        + "。\n\n## 子节\n\n- 一项\n- 二项\n\n1. 第一步\n2. 第二步\n\n| 名称 | 数量 |\n| --- | --- |\n| 中文 | 42 |\n\n"
        + tick * 3
        + 'python\nprint("你好")\n'
        + tick * 3
        + "\n"
    )
    raw = raw.replace("\\n", "\n")
    material = ingest_text(raw, name="新材料.md", material_id="m1", markdown=True)
    assert [block.kind for block in material.blocks] == [
        "heading",
        "para",
        "heading",
        "bullets",
        "numbered",
        "table",
        "para",
    ]
    assert [material.blocks[0].level, material.blocks[2].level] == [1, 2]
    assert material.blocks[1].text == "正文 加粗 与 代码。"
    assert material.blocks[3].items == ["一项", "二项"]
    assert material.blocks[5].columns == ["名称", "数量"]
    assert material.blocks[5].rows == [["中文", "42"]]
    assert 'print("你好")' in material.blocks[-1].text
    assert material.blocks[0].locator == "line:1-1"
    assert material.digest == hashlib.sha256(raw.encode()).hexdigest()
    assert (
        material.model_dump()
        == ingest_text(raw, name="新材料.md", material_id="m1", markdown=True).model_dump()
    )
    assert len({block.block_id for block in material.blocks}) == len(material.blocks)


def test_plain_text_never_interprets_markdown_and_digest_covers_whitespace() -> None:
    text = "# 原样" + chr(10) + "- 保留" + chr(10) * 2 + "下一段"
    material = ingest_text(text, name="a.txt", material_id="m1")
    assert [(b.kind, b.text) for b in material.blocks] == [
        ("para", "# 原样" + chr(10) + "- 保留"),
        ("para", "下一段"),
    ]
    changed = ingest_text(text + chr(10), name="a.txt", material_id="m1")
    assert changed.digest != material.digest


def test_markdown_external_images_html_and_nested_lists_remain_visible() -> None:
    text = chr(10).join(
        [
            "- 父项",
            "  - 子项",
            "",
            "![图示](https://invalid.example/p.png)",
            "",
            "<div>保留HTML文字</div>",
        ]
    )
    material = ingest_text(text, name="a.md", material_id="m1", markdown=True)
    content = chr(10).join(b.text + chr(10).join(b.items) for b in material.blocks)
    assert "父项" in content and "子项" in content
    assert "图示" in content and "https://invalid.example/p.png" in content
    assert "保留HTML文字" in content
    assert len(material.pending) >= 2


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("正文标题", 1)
    document.add_paragraph("正文内容")
    document.add_paragraph("无序项", style="List Bullet")
    document.add_paragraph("顺序项", style="List Number")
    table = document.add_table(rows=2, cols=2)
    for row, cells in zip(table.rows, [["名称", "数量"], ["测试", "7"]], strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="
    )
    document.add_picture(BytesIO(png))
    paragraph = document.add_paragraph("对象前文字")
    paragraph._p.append(OxmlElement("w:object"))
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_body_structure_images_and_unsupported_are_preserved() -> None:
    payload = _docx_bytes()
    material, assets = ingest_material(payload, name="a.docx", material_id="m1")
    assert [b.kind for b in material.blocks] == [
        "heading",
        "para",
        "bullets",
        "numbered",
        "table",
        "figure",
        "para",
    ]
    assert material.blocks[0].level == 1
    assert material.blocks[4].rows == [["测试", "7"]]
    assert material.blocks[5].image_ref in assets
    assert all(".." not in key and not key.startswith("/") for key in assets)
    assert next(iter(assets.values())).startswith(bytes([137, 80, 78, 71]))
    assert material.pending and "object" in str(material.pending)
    assert material.digest == hashlib.sha256(payload).hexdigest()
    repeated, repeated_assets = ingest_material(payload, name="a.docx", material_id="m1")
    assert repeated.model_dump() == material.model_dump() and repeated_assets == assets


def test_upload_text_encoding_suffix_and_empty_fail_explicitly() -> None:
    material, assets = ingest_material("正文".encode(), name="a.txt", material_id="m1")
    assert material.blocks[0].text == "正文" and not assets
    for payload, name in [
        (bytes([255]), "a.txt"),
        (b"abc", "a.pdf"),
        (b"bad", "a.docx"),
        (b" ", "a.txt"),
    ]:
        with pytest.raises(EngineError):
            ingest_material(payload, name=name, material_id="m1")


@pytest.mark.parametrize(
    "entry", ["../evil", "word/../evil", "/evil", "C:/evil", "word" + chr(92) + "evil"]
)
def test_docx_rejects_unsafe_package_paths(entry: str) -> None:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr(entry, b"bad")
    with pytest.raises(EngineError):
        ingest_material(buffer.getvalue(), name="a.docx", material_id="m1")


def test_docx_external_relationship_is_not_fetched_or_silently_lost() -> None:
    buffer = BytesIO()
    with ZipFile(BytesIO(_docx_bytes())) as source, ZipFile(buffer, "w", ZIP_DEFLATED) as output:
        for name in source.namelist():
            data = source.read(name)
            if name == "word/_rels/document.xml.rels":
                data = data.replace(
                    b'Target="media/image1.png"',
                    b'Target="https://invalid.example/p.png" TargetMode="External"',
                )
            output.writestr(name, data)
    material, assets = ingest_material(buffer.getvalue(), name="a.docx", material_id="m1")
    assert not assets
    assert any("外部" in pending for pending in material.pending)


def test_docx_unknown_drawing_and_content_controls_are_pending_not_lost() -> None:
    document = Document()
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "控件内文字"
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    control.append(content)
    document._element.body.insert(0, control)
    drawing = document.add_paragraph("图形旁文字").add_run()._r
    shape = OxmlElement("w:drawing")
    shape.append(OxmlElement("a:graphic"))
    drawing.append(shape)
    buffer = BytesIO()
    document.save(buffer)
    material, _ = ingest_material(buffer.getvalue(), name="a.docx", material_id="m1")
    assert any("控件内文字" in block.text for block in material.blocks)
    assert any("sdt" in pending for pending in material.pending)
    assert any("drawing" in pending for pending in material.pending)


def test_docx_equation_text_is_preserved_and_marked_pending() -> None:
    document = Document()
    paragraph = document.add_paragraph("公式：")
    equation = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "x=42"
    run.append(text)
    equation.append(run)
    paragraph._p.append(equation)
    buffer = BytesIO()
    document.save(buffer)
    material, _ = ingest_material(buffer.getvalue(), name="a.docx", material_id="m1")
    assert "x=42" in material.blocks[0].text
    assert any("oMath" in pending for pending in material.pending)


def test_material_unicode_surrogates_raise_engine_error() -> None:
    with pytest.raises(EngineError):
        ingest_text(chr(0xD800), name="a.txt", material_id="m1")

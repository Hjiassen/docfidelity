"""有效模板候选、选择与冻结的行为约束。"""

import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

import pytest

from docfidelity.contracts.errors import EngineError
from docfidelity.contracts.imports import SelectionRange, VariableMapping
from docfidelity.ooxml.doc import parse, qn, xpath
from docfidelity.ooxml.package import DocxPackage
from docfidelity.registry.candidates import inspect_template
from docfidelity.registry.freeze import freeze_template, load_template_version
from docfidelity.registry.selection import default_selection
from tests.fixtures.effective_template import build_effective_source

pytestmark = pytest.mark.unit


def test_freeze_erases_all_old_content_and_keeps_source(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    draft = inspect_template(source, import_id="draft")
    selection = default_selection(draft)
    frozen = freeze_template(
        source, selection, tmp_path / "bundle", template_id="t", version=1, name="T"
    )
    assert frozen.source_sha256 == before
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    with ZipFile(tmp_path / "bundle/effective.docx") as archive:
        for name in archive.namelist():
            assert b"OLD_" not in archive.read(name), name
        assert not any(
            "comments" in name or "embeddings" in name or "docProps" in name or "media" in name
            for name in archive.namelist()
        )
        assert b"PAGE" in archive.read("word/document.xml")
        assert b"INCLUDETEXT" not in archive.read("word/document.xml")
    for prototype in frozen.prototypes:
        assert b"OLD_" not in (tmp_path / "bundle" / prototype.path).read_bytes()
    assert frozen.format_bindings.body is not None
    assert frozen.format_bindings.heading
    assert len(frozen.format_bindings.tables) == 2
    assert load_template_version(tmp_path / "bundle") == frozen
    assert set(p.name for p in (tmp_path / "bundle").iterdir()) == {
        "effective.docx",
        "selection.json",
        "contract.json",
        "prototypes",
        "cleanup-report.json",
    }


def test_no_cover_is_unrecognized_and_manual_selection_changes_hash(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx", cover=False)
    draft = inspect_template(source, import_id="draft")
    assert inspect_template(source, import_id="draft").nodes == draft.nodes
    assert next(c for c in draft.candidates if c.kind == "cover").status == "unrecognized"
    selection = default_selection(draft)
    assert all(r.kind != "cover" for r in selection.ranges)
    first = freeze_template(
        source, selection, tmp_path / "one", template_id="t", version=1, name="T"
    )
    selection.ranges = [r for r in selection.ranges if r.kind != "header"]
    second = freeze_template(
        source, selection, tmp_path / "two", template_id="t", version=2, name="T"
    )
    assert first.content_hash != second.content_hash
    with ZipFile(tmp_path / "two/effective.docx") as archive:
        assert not any("header" in n for n in archive.namelist())
    assert json.loads((tmp_path / "two/selection.json").read_text("utf-8"))["ranges"]


@pytest.mark.parametrize("payload", [b"", b"broken"])
def test_bad_source_is_explicit_error(tmp_path: Path, payload: bytes) -> None:
    source = tmp_path / "bad.docx"
    source.write_bytes(payload)
    with pytest.raises(EngineError, match="E_TEMPLATE_UNREADABLE"):
        inspect_template(source, import_id="draft")


def test_variables_preserve_run_format_in_all_supported_containers(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    draft = inspect_template(source, import_id="draft")
    selection = default_selection(draft)
    previews = {
        "OLD_COVER": "title",
        "OLD_HEADER": "heading",
        "OLD_BODY": "value",
        "OLD_TABLE_ONE": "cell",
        "OLD_TEXTBOX": "box",
    }
    for text, key in previews.items():
        matching = [n for n in draft.nodes if n.kind == "paragraph" and n.text_preview == text]
        node = matching[-1]  # 文本框内层段落，而非同时包含它的外层段落。
        selection.variables.append(VariableMapping(key=key, node_id=node.node_id))
    frozen = freeze_template(
        source, selection, tmp_path / "bundle", template_id="t", version=1, name="T"
    )
    package = DocxPackage.open(tmp_path / "bundle/effective.docx")
    for variable in frozen.variables:
        matches = xpath(
            package.part(variable.part).read_tree,
            f'.//w:sdt[w:sdtPr/w:tag[@w:val="{variable.marker}"]]',
        )
        assert len(matches) == 1
        assert xpath(matches[0], "./w:sdtContent/w:p")
        if variable.key == "value":
            assert xpath(matches[0], ".//w:rPr/w:b")
            assert xpath(matches[0], ".//w:rPr/w:rFonts")[0].get(qn("w:ascii")) == "Arial"
    tables = [p for p in frozen.prototypes if p.kind == "table"]
    assert len(tables) == 2
    first, second = (parse((tmp_path / "bundle" / p.row_path).read_bytes()) for p in tables)
    assert xpath(first, ".//w:i")
    assert xpath(second, ".//w:i")[0].get(qn("w:val")) == "0"
    assert tables[0].columns == 2


@pytest.mark.parametrize(
    "case",
    [
        "out_of_range",
        "overlap",
        "key",
        "duplicate",
        "same_container",
        "unselected",
        "unsupported",
        "wrong_role",
    ],
)
def test_invalid_selection_is_rejected(tmp_path: Path, case: str) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    draft = inspect_template(source, import_id="draft")
    selection = default_selection(draft)
    body = next(n for n in draft.nodes if n.text_preview == "OLD_BODY")
    if case == "out_of_range":
        selection.ranges[0].end_node_id = "missing"
    elif case == "overlap":
        selection.ranges.append(selection.ranges[0])
    elif case == "wrong_role":
        next(r for r in selection.ranges if r.kind == "header").kind = "footer"
    elif case == "unsupported":
        obj = next(n for n in draft.nodes if n.kind == "object")
        block = next(
            n
            for n in draft.nodes
            if n.top_level and n.part == obj.part and n.block_index == obj.block_index
        )
        selection.ranges.append(
            SelectionRange(kind="body", start_node_id=block.node_id, end_node_id=block.node_id)
        )
    else:
        selection.variables = [
            VariableMapping(key="bad-key" if case == "key" else "value", node_id=body.node_id)
        ]
        if case == "duplicate":
            selection.variables.append(
                VariableMapping(
                    key="value",
                    node_id=next(n for n in draft.nodes if n.text_preview == "OLD_HEADING").node_id,
                )
            )
        elif case == "same_container":
            selection.variables.append(VariableMapping(key="other", node_id=body.node_id))
        elif case == "unselected":
            selection.ranges = []
    with pytest.raises(EngineError, match="E_SCHEMA_INVALID"):
        freeze_template(
            source, selection, tmp_path / "bundle", template_id="t", version=1, name="T"
        )
    assert not (tmp_path / "bundle").exists()


def test_manual_range_and_fixed_image_whitelist(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx", cover=False)
    draft = inspect_template(source, import_id="draft")
    selection = default_selection(draft)
    paragraph = next(n for n in draft.nodes if n.text_preview == "OLD_BODY")
    for span in selection.ranges:
        if span.start_node_id == paragraph.node_id:
            span.kind = "cover"
    selection.fixed_assets = [next(n.node_id for n in draft.nodes if n.kind == "image")]
    frozen = freeze_template(
        source, selection, tmp_path / "bundle", template_id="t", version=1, name="T"
    )
    with ZipFile(tmp_path / "bundle/effective.docx") as archive:
        media = [n for n in archive.namelist() if n.startswith("word/media/")]
        assert len(media) == 1
        assert archive.read(media[0]) == b"OLD_IMAGE"
        document = archive.read("word/document.xml")
        assert b"OLD_ALT" not in document
        assert b'cx="123" cy="456"' in document
    assert frozen.cleanup_report.fixed_assets == selection.fixed_assets


def test_section_roles_and_inheritance_are_retained(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source, read_only=False)
    document = package.part("word/document.xml").tree
    for section in xpath(document, ".//w:sectPr"):
        from docfidelity.ooxml.doc import sub

        sub(section, "w:titlePg")
    package.save(source)
    selection = default_selection(inspect_template(source, import_id="draft"))
    freeze_template(source, selection, tmp_path / "bundle", template_id="t", version=1, name="T")
    target = DocxPackage.open(tmp_path / "bundle/effective.docx")
    sections = xpath(target.part("word/document.xml").read_tree, ".//w:sectPr")
    assert len(sections) == 2
    assert len(xpath(sections[0], "./w:headerReference")) == 1
    assert not xpath(sections[1], "./w:headerReference")
    assert all(xpath(section, "./w:titlePg") for section in sections)
    assert all(resolved.part is not None for resolved in target.resolve_rels("word/document.xml"))


def test_frozen_versions_are_immutable_and_verify_after_source_deletion(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    selection = default_selection(inspect_template(source, import_id="draft"))
    versions = [
        freeze_template(
            source, selection, tmp_path / str(i), template_id="t", version=i + 1, name="T"
        )
        for i in range(3)
    ]
    assert len({v.content_hash for v in versions}) == 1
    with pytest.raises(EngineError, match="E_VERSION_CONFLICT"):
        freeze_template(source, selection, tmp_path / "0", template_id="t", version=1, name="T")
    source.unlink()
    assert load_template_version(tmp_path / "0") == versions[0]
    path = tmp_path / "0" / versions[0].prototypes[0].path
    path.write_bytes(b"tampered")
    with pytest.raises(EngineError, match="E_PACKAGE_INCOMPLETE"):
        load_template_version(tmp_path / "0")


def test_body_blocks_only_supply_prototypes_and_absence_is_explicit(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx", cover=False)
    draft = inspect_template(source, import_id="draft")
    selection = default_selection(draft)
    assert not selection.confirmed_absent
    selection.confirmed_absent = ["cover"]
    frozen = freeze_template(
        source, selection, tmp_path / "bundle", template_id="t", version=1, name="T"
    )
    package = DocxPackage.open(tmp_path / "bundle/effective.docx")
    assert not xpath(package.part("word/document.xml").read_tree, ".//w:tbl")
    assert frozen.selection.confirmed_absent == ["cover"]
    assert frozen.format_bindings.body is not None
    body = next(p for p in frozen.prototypes if p.prototype_id == frozen.format_bindings.body)
    assert xpath(parse((tmp_path / "bundle" / body.path).read_bytes()), ".//w:rPr/w:b")


def test_complex_safe_page_fields_and_section_switches(tmp_path: Path) -> None:
    from docfidelity.ooxml.doc import new_tree, sub

    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source, read_only=False)
    root = package.part("word/document.xml").tree
    body = xpath(root, "./w:body")[0]
    paragraph = new_tree("w:p")
    sub(sub(paragraph, "w:r"), "w:fldChar", {"w:fldCharType": "begin"})
    sub(sub(paragraph, "w:r"), "w:instrText").text = " PAGE "
    sub(sub(paragraph, "w:r"), "w:fldChar", {"w:fldCharType": "separate"})
    sub(sub(paragraph, "w:r"), "w:t").text = "OLD_PAGE_CACHE"
    sub(sub(paragraph, "w:r"), "w:fldChar", {"w:fldCharType": "end"})
    body.insert(len(body) - 1, paragraph)
    package.save(source)
    selection = default_selection(inspect_template(source, import_id="draft"))
    freeze_template(source, selection, tmp_path / "bundle", template_id="t", version=1, name="T")
    result = DocxPackage.open(tmp_path / "bundle/effective.docx")
    assert (
        len(xpath(result.part("word/document.xml").read_tree, './/w:fldSimple[@w:instr="PAGE"]'))
        == 2
    )


def test_missing_or_external_fixed_assets_are_explicit_errors(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source, read_only=False)
    blip = xpath(package.part("word/document.xml").tree, ".//a:blip")[0]
    blip.set(qn("r:link"), "external")
    package.save(source)
    draft = inspect_template(source, import_id="draft")
    selection = default_selection(draft)
    selection.fixed_assets = [next(n.node_id for n in draft.nodes if n.kind == "image")]
    with pytest.raises(EngineError, match="E_SCHEMA_INVALID"):
        freeze_template(
            source, selection, tmp_path / "bundle", template_id="t", version=1, name="T"
        )


def test_unreferenced_assets_are_excluded(tmp_path: Path) -> None:
    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source, read_only=False)
    package.add_binary_part("word/media/unreferenced.png", "image/png", b"OLD_UNREFERENCED")
    package.save(source)
    selection = default_selection(inspect_template(source, import_id="draft"))
    freeze_template(source, selection, tmp_path / "bundle", template_id="t", version=1, name="T")
    with ZipFile(tmp_path / "bundle/effective.docx") as archive:
        assert not any(n.startswith("word/media/") for n in archive.namelist())


def test_effective_package_opens_as_a_docx(tmp_path: Path) -> None:
    from docx import Document

    source = build_effective_source(tmp_path / "source.docx")
    selection = default_selection(inspect_template(source, import_id="draft"))
    freeze_template(source, selection, tmp_path / "bundle", template_id="t", version=1, name="T")
    document = Document(tmp_path / "bundle/effective.docx")
    assert len(document.sections) == 2
    assert all(not p.text for p in document.paragraphs)


def test_body_content_control_index_and_toc_cache_are_cleaned(tmp_path: Path) -> None:
    from docfidelity.ooxml.doc import new_tree, sub

    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source, read_only=False)
    body = xpath(package.part("word/document.xml").tree, "./w:body")[0]
    wrapper = new_tree("w:sdt")
    sub(sub(wrapper, "w:sdtPr"), "w:tag", {"w:val": "OLD_CONTROL_METADATA"})
    content = sub(wrapper, "w:sdtContent")
    field = sub(sub(content, "w:p"), "w:fldSimple", {"w:instr": "TOC"})
    sub(sub(field, "w:r"), "w:t").text = "OLD_TOC_CACHE"
    body.insert(2, wrapper)
    package.save(source)
    draft = inspect_template(source, import_id="draft")
    toc = next(c for c in draft.candidates if c.kind == "toc")
    assert toc.status == "detected"
    selection = default_selection(draft)
    freeze_template(source, selection, tmp_path / "bundle", template_id="t", version=1, name="T")
    with ZipFile(tmp_path / "bundle/effective.docx") as archive:
        document = archive.read("word/document.xml")
        assert b"OLD_TOC_CACHE" not in document
        assert b"OLD_CONTROL_METADATA" not in document
        assert b"TOC" not in document


def test_unsupported_drawing_is_skipped_until_selected(tmp_path: Path) -> None:
    from docfidelity.ooxml.doc import new_tree, sub

    source = build_effective_source(tmp_path / "source.docx")
    package = DocxPackage.open(source, read_only=False)
    body = xpath(package.part("word/document.xml").tree, "./w:body")[0]
    paragraph = new_tree("w:p")
    sub(sub(paragraph, "w:r"), "w:drawing")
    body.insert(len(body) - 1, paragraph)
    package.save(source)
    draft = inspect_template(source, import_id="draft")
    drawing = next(n for n in reversed(draft.nodes) if n.kind == "image")
    assert not drawing.supported
    selection = default_selection(draft)
    freeze_template(source, selection, tmp_path / "bundle", template_id="t", version=1, name="T")
    block = next(
        n
        for n in draft.nodes
        if n.top_level and n.part == drawing.part and n.block_index == drawing.block_index
    )
    selection.ranges.append(
        SelectionRange(kind="body", start_node_id=block.node_id, end_node_id=block.node_id)
    )
    with pytest.raises(EngineError, match="E_SCHEMA_INVALID"):
        freeze_template(source, selection, tmp_path / "bad", template_id="t", version=2, name="T")

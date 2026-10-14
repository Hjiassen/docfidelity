"""有效模板的合成夹具，不使用上传材料。"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START

from docfidelity.ooxml.doc import find_first, new_tree, qn, sub
from docfidelity.ooxml.package import DocxPackage


def build_effective_source(path: Path, *, cover: bool = True) -> Path:
    document = Document()
    if cover:
        document.add_paragraph("OLD_COVER", "Title")
        document.add_section(WD_SECTION_START.NEW_PAGE)
    document.sections[0].header.paragraphs[0].text = "OLD_HEADER"
    document.sections[0].footer.paragraphs[0].text = "OLD_FOOTER"
    document.add_heading("OLD_HEADING", 1)
    paragraph = document.add_paragraph()
    run = paragraph.add_run("OLD_BODY")
    run.bold = True
    run.font.name = "Arial"
    document.add_paragraph("OLD_LIST", "List Bullet")
    for label in ("ONE", "TWO"):
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = f"OLD_TABLE_{label}"
        table.cell(1, 0).paragraphs[0].add_run(f"OLD_ROW_{label}").italic = label == "ONE"
    document.core_properties.title = "OLD_PROPERTY"
    document.save(path)
    package = DocxPackage.open(path, read_only=False)
    body = find_first(package.part("word/document.xml").tree, "w:body")
    assert body is not None
    box = new_tree("w:p")
    content = sub(sub(sub(box, "w:r"), "w:pict"), "w:txbxContent")
    sub(sub(sub(content, "w:p"), "w:r"), "w:t").text = "OLD_TEXTBOX"
    body.insert(len(body) - 1, box)
    obj = new_tree("w:p")
    sub(sub(obj, "w:r"), "w:object").set("old", "OLD_OBJECT")
    body.insert(len(body) - 1, obj)
    field = new_tree("w:p")
    instruction = sub(field, "w:fldSimple", {"w:instr": 'INCLUDETEXT "https://OLD_EXTERNAL"'})
    sub(sub(instruction, "w:r"), "w:t").text = "OLD_FIELD"
    body.insert(len(body) - 1, field)
    page = new_tree("w:p")
    sub(page, "w:fldSimple", {"w:instr": "PAGE"})
    body.insert(len(body) - 1, page)
    comments = new_tree("w:comments")
    sub(sub(sub(sub(comments, "w:comment"), "w:p"), "w:r"), "w:t").text = "OLD_COMMENT"
    package.add_xml_part("word/comments.xml", "application/xml", comments)
    package.add_binary_part("word/embeddings/old.bin", "application/octet-stream", b"OLD_EMBEDDED")
    package.add_binary_part("word/media/old.png", "image/png", b"OLD_IMAGE")
    package.attach_relationship(
        "word/document.xml", type="http://example/relationships/image", target="media/old.png"
    )
    package.attach_relationship(
        "word/document.xml",
        type="http://example/relationships/hyperlink",
        target="https://OLD_LINK",
        target_mode="External",
    )
    # 绘图关联必须出现在文字容器中，才能测试明确确认素材。
    drawing = new_tree("w:p")
    inline = sub(sub(sub(drawing, "w:r"), "w:drawing"), "wp:inline")
    sub(inline, "wp:extent").attrib.update({"cx": "123", "cy": "456"})
    sub(inline, "wp:docPr").attrib.update({"id": "1", "name": "OLD_ALT", "descr": "OLD_ALT"})
    rel_id = package.relationships("word/document.xml")[-2].rel_id
    sub(sub(sub(inline, "a:graphic"), "a:graphicData"), "a:blip").set(qn("r:embed"), rel_id)
    body.insert(len(body) - 1, drawing)
    return package.save(path)

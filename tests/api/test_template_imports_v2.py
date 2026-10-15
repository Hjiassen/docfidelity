"""V2 导入提交、恢复和不可变版本的 HTTP 回归。"""

import base64
import multiprocessing
import warnings
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZipInfo

import pytest
from fastapi.testclient import TestClient
from PIL import Image

from server.main import create_app
from server.settings import ServerSettings
from tests.fixtures.effective_template import build_effective_source

pytestmark = pytest.mark.api


def client_at(root: Path) -> TestClient:
    return TestClient(
        create_app(
            settings=ServerSettings(
                storage_root=root, model_enabled=False, seed_demo_templates=False
            ),
            start_worker=False,
        )
    )


@pytest.fixture
def source(tmp_path: Path) -> bytes:
    return build_effective_source(tmp_path / "source.docx", cover=False).read_bytes()


def upload(client: TestClient, source: bytes) -> dict:
    response = client.post("/api/v2/template-imports", files={"file": ("format.docx", source)})
    assert response.status_code == 201, response.text
    assert response.headers["etag"] == f'"{response.json()["revision"]}"'
    return response.json()


def save(client: TestClient, draft: dict, selection: dict, *, match: str | None = None):
    return client.put(
        f'/api/v2/template-imports/{draft["import_id"]}/selection',
        json=selection,
        headers={"If-Match": match or str(draft["revision"])},
    )


def finalize(client: TestClient, draft: dict, name: str = "Format"):
    return client.post(
        f'/api/v2/template-imports/{draft["import_id"]}/finalize',
        json={"name": name},
        headers={"If-Match": str(draft["revision"])},
    )


def test_selection_survives_app_rebuild_and_conflicts(tmp_path: Path, source: bytes) -> None:
    root = tmp_path / "store"
    client = client_at(root)
    draft = upload(client, source)
    selection = draft["selection"]
    selection["ranges"] = [r for r in selection["ranges"] if r["kind"] != "header"]
    selection["confirmed_absent"] = ["header"]
    selection["formats"] = dict.fromkeys(
        ["page_setup", "heading", "body", "numbering", "tables"], False
    )
    body = next(n for n in draft["draft"]["nodes"] if n["text_preview"] == "OLD_BODY")
    selection["variables"] = [{"key": "title", "node_id": body["node_id"], "required": True}]
    updated = save(client, draft, selection, match=f'"{draft["revision"]}"')
    assert updated.status_code == 200, updated.text
    restored = client_at(root).get(f'/api/v2/template-imports/{draft["import_id"]}')
    assert restored.json()["selection"] == updated.json()["selection"]
    assert save(client, draft, selection).status_code == 409
    assert (
        client_at(root).get(f'/api/v2/template-imports/{draft["import_id"]}').json()["revision"]
        == 2
    )


def test_finalize_versions_are_immutable_and_render_download(tmp_path: Path, source: bytes) -> None:
    root = tmp_path / "store"
    client = client_at(root)
    draft = upload(client, source)
    frozen = finalize(client, draft)
    assert frozen.status_code == 201, frozen.text
    version = frozen.json()
    ref = f'{version["template_id"]}@{version["version"]}'
    bundle = root / "v2" / "templates" / version["template_id"] / "versions" / "1"
    old_contract = (bundle / "contract.json").read_bytes()
    assert finalize(client, draft).status_code == 409
    restored = client_at(root)
    detail = restored.get(f'/api/v2/templates/{version["template_id"]}')
    assert detail.json()["current_version"] == 1
    assert len(restored.get("/api/v2/templates").json()) == 1
    fresh = restored.get(f'/api/v2/template-imports/{draft["import_id"]}').json()
    assert finalize(restored, fresh, "Second").json()["version"] == 2
    assert (bundle / "contract.json").read_bytes() == old_contract
    # 渲染时原件可以不存在。
    (root / "v2" / "imports" / draft["import_id"] / "source.docx").unlink()
    rendered = restored.post(
        f'/api/v2/templates/{version["template_id"]}/versions/1/render',
        json={
            "content": {
                "project_id": "p",
                "doc_key": "d",
                "template_ref": ref,
                "units": [{"unit_id": "p", "type": "para", "text": "Fresh"}],
            },
            "variables": {},
        },
    )
    assert rendered.status_code == 200, rendered.text
    assert "attachment" in rendered.headers["content-disposition"]
    with ZipFile(BytesIO(rendered.content)) as archive:
        assert b"Fresh" in archive.read("word/document.xml")
    assert not list((root / "v2" / "renders").iterdir())


def test_concurrent_selection_has_one_winner(tmp_path: Path, source: bytes) -> None:
    root = tmp_path / "store"
    draft = upload(client_at(root), source)

    def put(index: int) -> int:
        selection = dict(draft["selection"], formats={"body": bool(index)})
        return save(client_at(root), draft, selection).status_code

    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(put, [0, 1])) == [200, 409]


@pytest.mark.parametrize(
    "member", ["../evil", "/evil", "C:/evil", "word//evil", "word/document.xml"]
)
def test_unsafe_zip_is_rejected(tmp_path: Path, source: bytes, member: str) -> None:
    payload = BytesIO(source)
    with ZipFile(payload, "a") as archive:
        info = ZipInfo("safe")
        info.filename = member
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Duplicate name:", category=UserWarning)
            archive.writestr(info, b"x")
    response = client_at(tmp_path / "store").post(
        "/api/v2/template-imports", files={"file": ("x.docx", payload.getvalue())}
    )
    assert response.status_code in {400, 415, 422}


def test_unknown_identifiers_return_404(tmp_path: Path) -> None:
    client = client_at(tmp_path / "store")
    for url in [
        "/api/v2/template-imports/i_missing",
        "/api/v2/templates/t_missing",
        "/api/v2/templates/t_missing/versions/1",
    ]:
        assert client.get(url).status_code == 404


def test_render_openapi_references_resolve(tmp_path: Path) -> None:
    spec = client_at(tmp_path / "store").get("/openapi.json").json()
    schema = spec["paths"]["/api/v2/templates/{template_id}/versions/{version}/render"]["post"][
        "requestBody"
    ]["content"]["application/json"]["schema"]
    assert schema == {"$ref": "#/components/schemas/RenderRequest"}
    assert "content" in spec["components"]["schemas"]["RenderRequest"]["properties"]

    def resolve(value):
        if isinstance(value, dict):
            if "$ref" in value:
                target = spec
                for segment in value["$ref"][2:].split("/"):
                    target = target[segment.replace("~1", "/").replace("~0", "~")]
                assert isinstance(target, dict)
            for child in value.values():
                resolve(child)
        elif isinstance(value, list):
            for child in value:
                resolve(child)

    resolve(spec["components"]["schemas"]["RenderRequest"])


def test_freeze_failure_preserves_selection(tmp_path: Path, source: bytes, monkeypatch) -> None:
    from docfidelity.contracts.errors import EngineError, ErrorCode
    from server.template_imports import engine

    root = tmp_path / "store"
    client = client_at(root)
    draft = upload(client, source)

    def fail(*args, **kwargs):
        raise EngineError(ErrorCode.SCHEMA_INVALID, "冻结失败注入")

    with monkeypatch.context() as patch:
        patch.setattr(engine, "freeze_template", fail)
        assert finalize(client, draft).status_code == 422
    restored = client_at(root)
    assert restored.get(f'/api/v2/template-imports/{draft["import_id"]}').json() == draft
    assert restored.get("/api/v2/templates").json() == []
    assert finalize(restored, draft).status_code == 201


def test_cache_failure_after_current_commit_still_returns_saved(
    tmp_path: Path, source: bytes, monkeypatch
) -> None:
    from server.template_imports import TemplateImports

    root = tmp_path / "store"
    client = client_at(root)
    draft = upload(client, source)

    def fail(*args, **kwargs):
        raise OSError("缓存写失败注入")

    with monkeypatch.context() as patch:
        patch.setattr(TemplateImports, "_cache", fail)
        response = finalize(client, draft)
        assert response.status_code == 201, response.text
    restored = client_at(root)
    assert finalize(restored, draft).status_code == 409
    version = response.json()
    assert (
        restored.get(f'/api/v2/templates/{version["template_id"]}').json()["current_version"] == 1
    )
    assert (root / "v2" / "templates" / version["template_id"] / "template.json").exists()


@pytest.mark.parametrize("operation", ["selection", "finalize"])
def test_current_failure_leaves_uncommitted_snapshots_invisible(
    tmp_path: Path, source: bytes, monkeypatch, operation: str
) -> None:
    import server.template_imports as storage
    from docfidelity.contracts.errors import EngineError, ErrorCode

    root = tmp_path / "store"
    client = client_at(root)
    draft = upload(client, source)
    original = storage.atomic_write_text

    def fail_pointer(path, text, **kwargs):
        if path.name == "current.json":
            raise EngineError(ErrorCode.INTERNAL, "current 写失败注入")
        return original(path, text, **kwargs)

    with monkeypatch.context() as patch:
        patch.setattr(storage, "atomic_write_text", fail_pointer)
        failed = (
            save(client, draft, draft["selection"])
            if operation == "selection"
            else finalize(client, draft)
        )
        assert failed.status_code == 500
    restored = client_at(root)
    assert restored.get(f'/api/v2/template-imports/{draft["import_id"]}').json() == draft
    assert restored.get("/api/v2/templates").json() == []
    if operation == "finalize":
        assert (
            restored.get(f'/api/v2/templates/{draft["template_id"]}/versions/1').status_code == 404
        )
        assert finalize(restored, draft).json()["version"] == 2
    else:
        assert save(restored, draft, draft["selection"]).json()["revision"] == 3


@pytest.mark.parametrize(
    "match", [None, 'W/"1"', "*", "1,2", '"bad"', pytest.param("1" * 5000, id="oversized-number")]
)
def test_invalid_if_match_never_saves(tmp_path: Path, source: bytes, match) -> None:
    client = client_at(tmp_path / "store")
    draft = upload(client, source)
    headers = {} if match is None else {"If-Match": match}
    response = client.put(
        f'/api/v2/template-imports/{draft["import_id"]}/selection',
        json=draft["selection"],
        headers=headers,
    )
    assert response.status_code == 409


@pytest.mark.parametrize("kind", ["symlink", "macro", "macro_type", "bomb", "not_docx"])
def test_more_zip_boundaries(tmp_path: Path, source: bytes, kind: str, monkeypatch) -> None:
    import server.template_import_io as io

    payload = BytesIO(source if kind != "not_docx" else b"")
    with ZipFile(payload, "a") as archive:
        info = ZipInfo("word/extra.bin")
        if kind == "symlink":
            info.external_attr = 0o120777 << 16
        if kind == "macro":
            info.filename = "word/VBAPROJECT.BIN"
        archive.writestr(info, b"x" * 500)
    if kind == "macro_type":
        rewritten = BytesIO()
        with ZipFile(BytesIO(payload.getvalue())) as old, ZipFile(rewritten, "w") as new:
            for info in old.infolist():
                data = old.read(info)
                if info.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"wordprocessingml.document.main+xml",
                        b"wordprocessingml.document.macroEnabled.main+xml",
                    )
                new.writestr(info, data)
        payload = rewritten
    if kind == "bomb":
        monkeypatch.setattr(io, "MAX_ZIP_EXPANDED_BYTES", 100)
    response = client_at(tmp_path / "store").post(
        "/api/v2/template-imports", files={"file": ("x.docx", payload.getvalue())}
    )
    assert response.status_code in {400, 413, 415, 422}


@pytest.mark.parametrize(
    ("encoding", "macro_type", "entities", "expected"),
    [
        ("utf-16", "application/vnd.ms-word.document.macroEnabled.main+xml", False, 415),
        ("utf-16", "application/vnd.ms-office.vbaProject", False, 415),
        ("utf-8", "application/vnd.ms-word.document.macroEnabled.main+xml", True, 415),
        ("utf-16", None, False, 201),
        ("utf-8", None, True, 201),
    ],
)
def test_content_type_xml_encoding_and_entities(
    tmp_path: Path,
    source: bytes,
    encoding: str,
    macro_type: str | None,
    entities: bool,
    expected: int,
) -> None:
    from lxml import etree

    rewritten = BytesIO()
    with ZipFile(BytesIO(source)) as old, ZipFile(rewritten, "w") as new:
        for info in old.infolist():
            data = old.read(info)
            if info.filename == "[Content_Types].xml":
                root = etree.fromstring(data)
                main = next(item for item in root if item.get("PartName") == "/word/document.xml")
                if macro_type is not None:
                    main.set("ContentType", macro_type)
                data = etree.tostring(root, encoding=encoding, xml_declaration=True)
                if entities:
                    text = data.decode(encoding)
                    value = main.get("ContentType")
                    text = text.replace(value, "".join(f"&#{ord(char)};" for char in value))
                    data = text.encode(encoding)
            new.writestr(info, data)
    response = client_at(tmp_path / "store").post(
        "/api/v2/template-imports", files={"file": ("x.docx", rewritten.getvalue())}
    )
    assert response.status_code == expected, response.text


@pytest.mark.parametrize(
    "content_types", [b"<Types>", b'<!DOCTYPE Types [<!ENTITY kind "safe">]><Types>&kind;</Types>']
)
def test_content_types_malformed_or_dtd_is_rejected(
    tmp_path: Path,
    source: bytes,
    content_types: bytes,
) -> None:
    rewritten = BytesIO()
    with ZipFile(BytesIO(source)) as old, ZipFile(rewritten, "w") as new:
        for info in old.infolist():
            new.writestr(
                info, content_types if info.filename == "[Content_Types].xml" else old.read(info)
            )
    response = client_at(tmp_path / "store").post(
        "/api/v2/template-imports", files={"file": ("x.docx", rewritten.getvalue())}
    )
    assert response.status_code == 422, response.text


def _process_save(root: str, draft: dict, gate, queue) -> None:
    gate.wait(20)
    queue.put(save(client_at(Path(root)), draft, draft["selection"]).status_code)


def test_process_selection_has_one_winner(tmp_path: Path, source: bytes) -> None:
    root = tmp_path / "store"
    draft = upload(client_at(root), source)
    context = multiprocessing.get_context("spawn")
    gate, queue = context.Event(), context.Queue()
    processes = [
        context.Process(target=_process_save, args=(str(root), draft, gate, queue))
        for _ in range(2)
    ]
    try:
        for process in processes:
            process.start()
        gate.set()
        assert sorted(queue.get(timeout=30) for _ in processes) == [200, 409]
        for process in processes:
            process.join(10)
            assert process.exitcode == 0
    finally:
        for process in processes:
            if process.is_alive():
                process.terminate()
            process.join(10)


@pytest.mark.parametrize(
    "asset", ["good", "../outside.png", "C:/outside.png", "bad_base64", "bad_image", "missing"]
)
def test_render_image_assets_are_logical_and_validated(
    tmp_path: Path, source: bytes, asset: str
) -> None:
    root = tmp_path / "store"
    client = client_at(root)
    draft = upload(client, source)
    version = finalize(client, draft).json()
    image = BytesIO()
    Image.new("RGB", (4, 4), "red").save(image, format="PNG")
    encoded = base64.b64encode(image.getvalue()).decode("ascii")
    name = "image.png" if asset in {"good", "bad_base64", "bad_image", "missing"} else asset
    assets = {name: encoded}
    if asset == "bad_base64":
        assets[name] = "!!!"
    if asset == "bad_image":
        assets[name] = base64.b64encode(b"\x89PNG\r\n\x1a\n" + b"x" * 32).decode("ascii")
    if asset == "missing":
        assets = {}
    body = {
        "content": {
            "project_id": "p",
            "doc_key": "d",
            "template_ref": f'{version["template_id"]}@1',
            "units": [{"unit_id": "image", "type": "figure", "fields": {"image_ref": name}}],
        },
        "variables": {},
        "assets": assets,
    }
    response = client.post(
        f'/api/v2/templates/{version["template_id"]}/versions/1/render', json=body
    )
    assert response.status_code == (
        200
        if asset == "good"
        else 400
        if ":" in asset or ".." in asset
        else 415
        if asset in {"bad_base64", "bad_image"}
        else 422
    ), response.text
    if asset == "good":
        with ZipFile(BytesIO(response.content)) as archive:
            assert any(name.startswith("word/media/") for name in archive.namelist())
    assert not list((root / "v2" / "renders").iterdir())


def test_render_stream_limit_and_asset_limit(tmp_path: Path, source: bytes, monkeypatch) -> None:
    from server import template_import_io as io

    root = tmp_path / "store"
    client = client_at(root)
    version = finalize(client, upload(client, source)).json()
    url = f'/api/v2/templates/{version["template_id"]}/versions/1/render'
    bounded = TestClient(
        create_app(
            settings=ServerSettings(storage_root=root, max_upload_bytes=64), start_worker=False
        )
    )
    assert (
        bounded.post(
            url,
            content=b"x" * 65,
            headers={"content-type": "application/json", "content-length": "1"},
        ).status_code
        == 413
    )
    image = BytesIO()
    Image.new("RGB", (4, 4)).save(image, format="PNG")
    monkeypatch.setattr(io, "MAX_ASSET_BYTES", 2)
    response = client.post(
        url,
        json={
            "content": {
                "project_id": "p",
                "doc_key": "d",
                "template_ref": f'{version["template_id"]}@1',
            },
            "assets": {"image.png": base64.b64encode(image.getvalue()).decode("ascii")},
        },
    )
    assert response.status_code == 413
    assert not list((root / "v2" / "renders").iterdir())


@pytest.mark.parametrize("member", [b"word\\evil", b"word/ev\x00l"])
def test_raw_zip_path_syntax_is_not_normalized_away(
    tmp_path: Path, source: bytes, member: bytes
) -> None:
    payload = BytesIO(source)
    with ZipFile(payload, "a") as archive:
        archive.writestr("word/evil", b"x")
    raw = payload.getvalue().replace(b"word/evil", member)
    response = client_at(tmp_path / "store").post(
        "/api/v2/template-imports", files={"file": ("x.docx", raw)}
    )
    assert response.status_code == 400


def test_fixed_assets_and_manual_ranges_survive_rebuild(tmp_path: Path) -> None:
    from docx import Document

    image = BytesIO()
    Image.new("RGB", (4, 4), "red").save(image, format="PNG")
    image.seek(0)
    document = Document()
    document.add_heading("Source heading", 1)
    document.add_paragraph("Source body")
    document.add_picture(image)
    data = BytesIO()
    document.save(data)
    root = tmp_path / "store"
    client = client_at(root)
    draft = upload(client, data.getvalue())
    selection = draft["selection"]
    node = next(n for n in draft["draft"]["nodes"] if n["kind"] == "image")
    selection["fixed_assets"] = [node["node_id"]]
    body = next(n for n in draft["draft"]["nodes"] if n["text_preview"] == "Source body")
    selection["ranges"] = [r for r in selection["ranges"] if r["start_node_id"] != body["node_id"]]
    selection["ranges"].append(
        {"kind": "body", "start_node_id": body["node_id"], "end_node_id": body["node_id"]}
    )
    saved = save(client, draft, selection)
    assert saved.status_code == 200, saved.text
    restored = client_at(root).get(f'/api/v2/template-imports/{draft["import_id"]}').json()
    assert restored["selection"] == saved.json()["selection"]
    version = finalize(client_at(root), restored)
    assert version.status_code == 201, version.text
    assert version.json()["selection"]["fixed_assets"] == [node["node_id"]]


def test_concurrent_finalize_has_one_committed_version(tmp_path: Path, source: bytes) -> None:
    root = tmp_path / "store"
    draft = upload(client_at(root), source)

    def freeze(index: int) -> int:
        return finalize(client_at(root), draft, f"Format {index}").status_code

    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(freeze, [0, 1])) == [201, 409]
    detail = client_at(root).get(f'/api/v2/templates/{draft["template_id"]}').json()
    assert detail["versions"] == [1]


@pytest.mark.parametrize("setting", ["MAX_ZIP_MEMBERS", "MAX_ZIP_MEMBER_BYTES", "MAX_ZIP_RATIO"])
def test_zip_resource_limits(tmp_path: Path, source: bytes, setting: str, monkeypatch) -> None:
    import server.template_import_io as io

    monkeypatch.setattr(io, setting, 1)
    response = client_at(tmp_path / "store").post(
        "/api/v2/template-imports", files={"file": ("x.docx", source)}
    )
    assert response.status_code == 413


def test_bad_variable_render_cleans_temporary_directory(tmp_path: Path, source: bytes) -> None:
    root = tmp_path / "store"
    client = client_at(root)
    draft = upload(client, source)
    version = finalize(client, draft).json()
    response = client.post(
        f'/api/v2/templates/{version["template_id"]}/versions/1/render',
        json={
            "content": {
                "project_id": "p",
                "doc_key": "d",
                "template_ref": f'{version["template_id"]}@1',
            },
            "variables": {"unknown": "value"},
        },
    )
    assert response.status_code == 422
    assert not list((root / "v2" / "renders").iterdir())

"""V2 文档、事务与运行的真实 HTTP 回归。"""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient

from server.main import create_app
from server.settings import ServerSettings
from tests.fixtures.effective_template import build_effective_source

pytestmark = pytest.mark.api


def client_at(root: Path, **kwargs):
    return TestClient(
        create_app(
            settings=ServerSettings(
                _env_file=None, storage_root=root, model_enabled=False, seed_demo_templates=False
            ),
            start_worker=False,
            **kwargs,
        )
    )


def setup_document(tmp_path: Path, *, mode="format_only", **kwargs):
    client = client_at(tmp_path / "store", **kwargs)
    source = build_effective_source(tmp_path / "source.docx", cover=False).read_bytes()
    imported = client.post(
        "/api/v2/template-imports", files={"file": ("format.docx", source)}
    ).json()
    version = client.post(
        f'/api/v2/template-imports/{imported["import_id"]}/finalize',
        json={"name": "Format"},
        headers={"If-Match": "1"},
    ).json()
    material = client.post(
        "/api/v2/materials",
        json={"name": "new.md", "text": "# 新章\n新材料 42。", "markdown": True},
    )
    assert material.status_code == 201, material.text
    document = client.post(
        "/api/v2/documents",
        json={
            "name": "新文档",
            "template_id": version["template_id"],
            "template_version": 1,
            "material_ids": [material.json()["material_id"]],
            "mode": mode,
        },
    )
    assert document.status_code == 201, document.text
    assert document.headers["etag"] == '"1"'
    return client, document.json(), material.json()


def submit(client, document, kind, **kwargs):
    response = client.post(
        f'/api/v2/documents/{document["document_id"]}/{kind}-runs', json={}, **kwargs
    )
    assert response.status_code == 202, response.text
    return response.json()


def drain(client):
    return client.app.state.docfidelity_state.document_runs.run_until_idle()


def test_material_revision_and_download_survive_restart(tmp_path: Path):
    client, document, material = setup_document(tmp_path)
    run = submit(client, document, "outline")
    assert run["state"] == "queued"
    drain(client)
    document = client.get(f'/api/v2/documents/{document["document_id"]}').json()
    assert document["outline"][0]["title"] == "新章"
    response = client.patch(
        f'/api/v2/documents/{document["document_id"]}',
        json={"outline_confirmed": True},
        headers={"If-Match": str(document["revision"])},
    )
    assert response.status_code == 200, response.text
    document = response.json()
    content = submit(client, document, "content")
    drain(client)
    finished = client.get(f'/api/v2/runs/{content["run_id"]}').json()
    assert finished["state"] == "succeeded"
    assert finished["attempts"] == []
    document = client.get(f'/api/v2/documents/{document["document_id"]}').json()
    assert document["sections"][0]["blocks"][0]["text"] == "新材料 42。"
    render = submit(client, document, "render")
    drain(client)
    restored = client_at(tmp_path / "store")
    artifacts = restored.get(f'/api/v2/runs/{render["run_id"]}/artifacts').json()
    assert len(artifacts) == 1
    download = restored.get(
        f'/api/v2/runs/{render["run_id"]}/artifacts/{artifacts[0]["artifact_id"]}'
    )
    assert download.status_code == 200
    with ZipFile(BytesIO(download.content)) as archive:
        assert "新材料 42。" in archive.read("word/document.xml").decode()
    assert restored.get(f'/api/v2/materials/{material["material_id"]}').json() == material
    original = restored.get(f'/api/v2/documents/{document["document_id"]}/revisions/1').json()
    assert original["outline"] == []


def test_etag_full_validation_idempotency_and_cancel(tmp_path: Path):
    client, document, _ = setup_document(tmp_path)
    path = f'/api/v2/documents/{document["document_id"]}'
    assert client.patch(path, json={"name": "changed"}).status_code == 409
    assert (
        client.patch(path, json={"outline_confirmed": True}, headers={"If-Match": "1"}).status_code
        == 422
    )
    first = submit(client, document, "outline", headers={"Idempotency-Key": "once"})
    assert (
        submit(client, document, "outline", headers={"Idempotency-Key": "once"})["run_id"]
        == first["run_id"]
    )
    assert client.post(f'/api/v2/runs/{first["run_id"]}/cancel').json()["state"] == "canceled"
    drain(client)
    assert client.get(path).json()["revision"] == 1
    assert client.get(f'/api/v2/runs/{first["run_id"]}/artifacts').json() == []


def test_uploaded_material_sources_and_reject_configuration_override(tmp_path: Path):
    client = client_at(tmp_path)
    uploaded = client.post(
        "/api/v2/materials/upload", files={"file": ("新.md", "# 章节\n新事实。".encode())}
    )
    assert uploaded.status_code == 201, uploaded.text
    assert uploaded.json()["blocks"][0]["locator"].startswith("line:")
    assert (
        client.post(
            "/api/v2/materials",
            json={"name": "a.txt", "text": "x", "model_base_url": "http://evil"},
        ).status_code
        == 422
    )


def test_failed_pointer_never_exposes_or_overwrites_revision(tmp_path: Path, monkeypatch):
    from server import document_store

    client, document, _ = setup_document(tmp_path)
    original = document_store.atomic_write_text

    def fail_pointer(path, text, **kwargs):
        if path.name == "current.json":
            raise OSError("simulated disk failure")
        return original(path, text, **kwargs)

    monkeypatch.setattr(document_store, "atomic_write_text", fail_pointer)
    response = client.patch(
        f'/api/v2/documents/{document["document_id"]}',
        json={"name": "invisible"},
        headers={"If-Match": "1"},
    )
    assert response.status_code == 500
    assert client.get(f'/api/v2/documents/{document["document_id"]}').json() == document
    assert client.get(f'/api/v2/documents/{document["document_id"]}/revisions/2').status_code == 404
    monkeypatch.setattr(document_store, "atomic_write_text", original)
    response = client.patch(
        f'/api/v2/documents/{document["document_id"]}',
        json={"name": "visible"},
        headers={"If-Match": "1"},
    )
    assert response.json()["revision"] == 3
    assert client.get(f'/api/v2/documents/{document["document_id"]}/revisions/2').status_code == 404


def test_same_key_with_changed_match_conflicts(tmp_path: Path):
    client, document, _ = setup_document(tmp_path)
    submit(client, document, "outline", headers={"If-Match": "1", "Idempotency-Key": "once"})
    response = client.post(
        f'/api/v2/documents/{document["document_id"]}/outline-runs',
        json={},
        headers={"If-Match": "2", "Idempotency-Key": "once"},
    )
    assert response.status_code == 409


def test_outline_change_invalidates_confirmation_and_stale_sections(tmp_path: Path):
    client, document, _ = setup_document(tmp_path)
    submit(client, document, "outline")
    drain(client)
    path = f'/api/v2/documents/{document["document_id"]}'
    document = client.get(path).json()
    document = client.patch(
        path, json={"outline_confirmed": True}, headers={"If-Match": str(document["revision"])}
    ).json()
    submit(client, document, "content")
    drain(client)
    document = client.get(path).json()
    outline = document["outline"]
    outline[0]["source_refs"] = []
    updated = client.patch(
        path, json={"outline": outline}, headers={"If-Match": str(document["revision"])}
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["outline_confirmed"] is False
    assert updated.json()["sections"] == []


def test_images_are_verified_content_bound_and_frozen(tmp_path: Path, monkeypatch):
    from docx import Document
    from PIL import Image

    from docfidelity.contracts.errors import EngineError

    image = BytesIO()
    Image.new("RGB", (8, 8), "red").save(image, format="PNG")
    source = BytesIO()
    doc = Document()
    doc.add_paragraph("新图片")
    doc.add_picture(BytesIO(image.getvalue()))
    doc.save(source)
    client = client_at(tmp_path)
    response = client.post(
        "/api/v2/materials/upload", files={"file": ("picture.docx", source.getvalue())}
    )
    assert response.status_code == 201, response.text
    material = response.json()
    service = client.app.state.docfidelity_state.materials
    stored = service.get(material["material_id"])
    assets = service.assets([stored])
    reference = next(iter(assets))
    (tmp_path / "v2" / "materials" / material["material_id"] / "assets" / reference).write_bytes(
        b"tampered"
    )
    with pytest.raises(EngineError, match="摘要"):
        service.assets([stored])


def test_json_text_keeps_name_and_explicit_plain_text_semantics(tmp_path: Path):
    client = client_at(tmp_path)
    result = client.post(
        "/api/v2/materials", json={"name": "notes.md", "text": "# literal", "markdown": False}
    )
    assert result.status_code == 201, result.text
    assert result.json()["name"] == "notes.md"
    assert result.json()["media_type"] == "text"
    assert result.json()["blocks"][0]["kind"] == "para"
